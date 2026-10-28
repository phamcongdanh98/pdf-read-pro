import hashlib
import os
import tempfile as _tf
os.environ["APPDATA"] = _tf.mkdtemp()      # thùng rác / dữ liệu app của test không lẫn vào dữ liệu thật
import random
import shutil
import sys
import tempfile
import unittest
from unittest import mock

import fitz
import pikepdf

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)
from pdfread_pro.core import compressor as C, pipeline as P, renamer as R, rules1361, scanner, signgather as G  # noqa: E402
from make_signed_sample import sign_pdf  # noqa: E402
from test_compressor import scan_page  # noqa: E402

A, B = "056123456789_TranVanThai", "079111222333_LeThiChuyen"


def tree_hash(root):
    """Dấu vân tay cả cây thư mục: (đường dẫn tương đối, kích thước, md5 nội dung)."""
    out = []
    for dp, _, fns in os.walk(root):
        for f in sorted(fns):
            p = os.path.join(dp, f)
            with open(p, "rb") as fh:
                out.append((os.path.relpath(p, root), os.path.getsize(p), hashlib.md5(fh.read()).hexdigest()))
    return sorted(out)


def text_pdf(path, pages=1):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = fitz.open()
    for i in range(pages):
        d.new_page().insert_text((72, 72), f"trang {i} {os.path.basename(path)}")
    d.save(path)
    d.close()


def heavy_pdf(path, pages=2):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    rnd = random.Random(3)
    d = fitz.open()
    for _ in range(pages):
        scan_page(d, rnd, 1240, 1754)
    d.save(path)
    d.close()


class PipelineTests(unittest.TestCase):
    def setUp(self):
        G.BUSY_SECONDS = 0
        self.tmp = tempfile.mkdtemp()
        self.src = os.path.join(self.tmp, "Ho so")
        a, b = os.path.join(self.src, "Chi bo 1", A), os.path.join(self.src, "Chi bo 1", B)
        heavy_pdf(os.path.join(a, "03 phieu dang vien.pdf"))                        # nặng → nén; tên có mã 03
        text_pdf(os.path.join(a, "1.ly lich.pdf"))                                  # mã '1' → 01.…
        text_pdf(os.path.join(a, "xyz_lung_tung.pdf"))                              # không có mã → giữ tên
        text_pdf(os.path.join(a, "tmp5.pdf"), 2)
        sign_pdf(os.path.join(a, "tmp5.pdf"), os.path.join(a, "05.QD_ket_nap_dang_vien.signed.pdf"), "TEST")
        os.remove(os.path.join(a, "tmp5.pdf"))
        text_pdf(os.path.join(b, "02 ly lich dang vien.pdf"), 3)
        with open(os.path.join(a, "ghi_chu.txt"), "w", encoding="utf-8") as fh:
            fh.write("không phải PDF")
        text_pdf(os.path.join(self.src, "Khac", "ghi_chu.pdf"))                     # ngoài thư mục CCCD
        self.work = os.path.join(self.tmp, "Ho so - Quy trinh")
        self.gather = self.work + " - Gom ky so"
        self.cfg = P.Config(self.src, self.work, self.gather, limit_kb=250, mode="1361",
                            rules=rules1361.default_rules(), rule_label="1361")
        self.before = tree_hash(self.src)

    def tearDown(self):
        G.BUSY_SECONDS = 5
        shutil.rmtree(self.tmp, ignore_errors=True)

    def files(self):
        return scanner.scan(self.src).files

    # ------------------------------------------------------------------ chạy trọn vẹn
    def test_full_pipeline(self):
        stages = []
        rep = P.run(self.cfg, P.Hooks(stage=lambda i, t: stages.append(i)))
        self.assertEqual(stages, [0, 1, 2, 3])
        self.assertEqual(rep.error, "")
        self.assertTrue(rep.ok, [(s.name, s.status, s.summary) for s in rep.stages])
        self.assertEqual(tree_hash(self.src), self.before)                          # HỒ SƠ GỐC KHÔNG ĐỔI

        wa = os.path.join(self.work, "Chi bo 1", A)
        names = sorted(os.listdir(wa))
        self.assertIn("01.Ly_lich_nguoi_xin_vao_Dang.pdf", names)                   # '1.ly lich' → 01.…
        self.assertIn("03.Phieu_dang_vien.pdf", names)
        self.assertIn("xyz_lung_tung.pdf", names)                                   # không có quy tắc → giữ nguyên
        self.assertIn("ghi_chu.txt", names)                                         # file khác vẫn được chép
        self.assertIn("02.Ly_lich_dang_vien.pdf", os.listdir(os.path.join(self.work, "Chi bo 1", B)))
        self.assertTrue(os.path.isfile(os.path.join(self.work, P.MANIFEST)))

        # nén: trang nặng ≤ 250 KB, đủ số trang; file đã ký giữ nguyên từng byte
        with pikepdf.open(os.path.join(wa, "03.Phieu_dang_vien.pdf")) as pdf:
            self.assertEqual(len(pdf.pages), 2)
            self.assertTrue(all(s <= 250 * 1024 for s in C.page_sizes(pdf).values()))
        signed_src = os.path.join(self.src, "Chi bo 1", A, "05.QD_ket_nap_dang_vien.signed.pdf")
        with open(signed_src, "rb") as f1, open(os.path.join(wa, "05.QD_ket_nap_dang_vien.signed.pdf"), "rb") as f2:
            self.assertEqual(f1.read(), f2.read())
        self.assertGreater(rep.saved_bytes, 0)

        # gom: chỉ file chưa ký trong thư mục CCCD_HoVaTen; không gom file đã ký / ngoài thư mục hồ sơ
        g = sorted(n for n in os.listdir(self.gather) if n.endswith(".pdf"))
        self.assertEqual(g, [f"01.Ly_lich_nguoi_xin_vao_Dang_{A}.pdf", f"02.Ly_lich_dang_vien_{B}.pdf",
                             f"03.Phieu_dang_vien_{A}.pdf", f"xyz_lung_tung_{A}.pdf"])
        self.assertTrue(os.path.isfile(os.path.join(self.gather, G.MANIFEST)))
        by_old = {r.old_name: r for r in rep.recs}
        self.assertEqual(by_old["ghi_chu.pdf"].gather_level, P.WARN)                # ngoài CCCD_HoVaTen
        self.assertEqual(by_old["xyz_lung_tung.pdf"].rename_level, P.WARN)
        self.assertEqual(by_old["1.ly lich.pdf"].new_name, "01.Ly_lich_nguoi_xin_vao_Dang.pdf")
        self.assertIn("Giữ nguyên: đã ký", by_old["05.QD_ket_nap_dang_vien.signed.pdf"].comp)
        self.assertEqual(rep.gathered, 4)
        # bàn giao cho tab Gom ký số: kiểm tra trả về nhận ra thư mục gom, trỏ đúng bản làm việc
        self.assertEqual(G.guess_root(self.gather), self.work)                      # sổ ghi thư mục gốc = bản làm việc
        self.assertTrue(G.looks_like_gather(self.gather))

    def test_each_step_only_touches_work_copy_even_on_rerun_guard(self):
        with self.assertRaises(P.PipelineError):
            P._guard(os.path.join(self.src, "Chi bo 1", A, "x.pdf"), self.work, self.gather)
        self.assertEqual(P._guard(os.path.join(self.work, "a.pdf"), self.work), os.path.join(self.work, "a.pdf"))

    # ------------------------------------------------------------------ kiểm tra trước
    def test_preflight_counts_and_problems(self):
        pf = P.preflight(self.cfg, self.files())
        self.assertEqual(pf.problems, [])
        self.assertEqual(pf.files, 6)
        self.assertEqual(pf.other_files, 1)                                          # ghi_chu.txt
        self.assertEqual(pf.rename_ready, 3)                  # '1.ly lich', '03 phieu…', '02 ly lich dang vien'
        self.assertEqual(pf.rename_norule, 2)                 # xyz_lung_tung.pdf, ghi_chu.pdf (không có mã)
        self.assertEqual(pf.rename_same, 1)                   # file .signed đã đúng tên
        self.assertEqual(pf.signed, 1)
        self.assertGreaterEqual(pf.est_compress, 1)
        self.assertEqual(pf.gather_new, 4)
        self.assertEqual(pf.gather_skip_signed, 0)
        self.assertEqual(pf.gather_not_dossier, 1)
        self.assertTrue(any("chưa có quy tắc" in w for w in pf.warnings))
        self.assertEqual(tree_hash(self.src), self.before)                           # kiểm tra trước không ghi gì
        self.assertFalse(os.path.exists(self.work))
        # thư mục sai
        bad = P.Config(self.src, os.path.join(self.src, "con"), self.gather, rules=self.cfg.rules)
        self.assertTrue(any("trùng / lồng" in p for p in P.preflight(bad, self.files()).problems))
        os.makedirs(self.work)
        with open(os.path.join(self.work, "cu.txt"), "w") as fh:
            fh.write("x")
        self.assertTrue(any("KHÔNG rỗng" in p for p in P.preflight(self.cfg, self.files()).problems))
        # ổ không đủ chỗ
        shutil.rmtree(self.work)
        with mock.patch.object(P, "_free_bytes", return_value=1000):
            self.assertTrue(any("không đủ chỗ" in p for p in P.preflight(self.cfg, self.files()).problems))
        # bộ quy tắc rỗng / giới hạn sai
        self.assertTrue(any("đang trống" in p for p in P.preflight(P.Config(self.src, self.work, self.gather),
                                                                   self.files()).problems))
        c2 = P.Config(self.src, self.work, self.gather, limit_kb=5, rules=self.cfg.rules)
        self.assertTrue(any("20 đến 20.000" in p for p in P.preflight(c2, self.files()).problems))

    def test_default_dirs_unique(self):
        w1, g1 = P.default_dirs(self.src)
        self.assertTrue(w1.endswith(" - Quy trình") and g1 == w1 + " - Gom ký số")
        os.makedirs(w1)
        w2, g2 = P.default_dirs(self.src)
        self.assertTrue(w2.endswith("(2)") and w2 != w1 and not os.path.exists(g2))

    # ------------------------------------------------------------------ lỗi / dừng / khôi phục
    def test_cancel_leaves_source_untouched_and_skips_later_steps(self):
        calls = {"n": 0}

        def cancelled():
            calls["n"] += 1
            return calls["n"] > 3000 or self._stage >= 1

        self._stage = 0
        rep = P.run(self.cfg, P.Hooks(stage=lambda i, t: setattr(self, "_stage", i), cancelled=cancelled))
        self.assertTrue(rep.cancelled)
        self.assertFalse(rep.ok)
        self.assertEqual(tree_hash(self.src), self.before)
        self.assertFalse(os.path.exists(self.gather))                                # chưa tới bước gom
        self.assertEqual(rep.stages[3].status, P.SKIP)

    def test_copy_failure_stops_everything(self):
        with mock.patch.object(G.shutil, "disk_usage", return_value=mock.Mock(free=10)):
            rep = P.run(self.cfg)
        self.assertFalse(rep.ok)
        self.assertIn("không đủ chỗ", rep.error)
        self.assertEqual(rep.stages[0].status, P.FAIL)
        self.assertTrue(all(s.status == P.SKIP for s in rep.stages[1:]))
        self.assertEqual(tree_hash(self.src), self.before)
        self.assertFalse(os.path.exists(self.gather))

    def test_run_refuses_unsafe_dirs(self):
        cfg = P.Config(self.src, os.path.join(self.src, "lam viec"), self.gather, rules=self.cfg.rules)
        rep = P.run(cfg)
        self.assertIn("trùng / lồng", rep.error)
        self.assertEqual(tree_hash(self.src), self.before)
        self.assertFalse(os.path.exists(os.path.join(self.src, "lam viec")))

    def test_broken_compressed_file_is_restored_from_original(self):
        os.makedirs(os.path.join(self.work, "Chi bo 1", A))
        rel = os.path.join("Chi bo 1", A, "03 phieu dang vien.pdf")
        shutil.copy2(os.path.join(self.src, rel), os.path.join(self.work, rel))
        path = os.path.join(self.work, rel)
        with open(path, "wb") as fh:
            fh.write(b"%PDF-1.4 hong")                                               # giả lập nén ra file hỏng
        rec = P.Rec(rel, A, "03 phieu dang vien.pdf", path, before=os.path.getsize(os.path.join(self.src, rel)))
        r = C.Result(path=path, status=C.ST_COMPRESSED, before=rec.before, after=13, pages=2)
        stats = {"ok": 0, "partial": 0, "fail": 0, "restored": 0, "before": 0, "after": 0}
        P._apply_compress(rec, r, self.cfg, stats)
        self.assertEqual(stats["restored"], 1)
        self.assertIn("khôi phục", rec.comp)
        self.assertEqual(rec.comp_level, P.WARN)
        self.assertEqual(os.path.getsize(path), rec.before)                          # đã trả về bản gốc
        with pikepdf.open(path) as pdf:
            self.assertEqual(len(pdf.pages), 2)
        self.assertEqual(tree_hash(self.src), self.before)

    def test_restore_blocked_outside_work(self):
        rec = P.Rec("a.pdf", A, "a.pdf", os.path.join(self.src, "Chi bo 1", A, "03 phieu dang vien.pdf"))
        with self.assertRaises(P.PipelineError):
            P._restore(rec, self.cfg)                                                # đích trong hồ sơ gốc → từ chối
        self.assertEqual(tree_hash(self.src), self.before)

    def test_custom_ruleset_mode(self):
        ruleset = R.default_presets()[0]
        cfg = P.Config(self.src, self.work, self.gather, mode="custom", ruleset=ruleset, policy="auto", rule_label="mẫu")
        pf = P.preflight(cfg, self.files())
        self.assertEqual(pf.problems, [])
        rep = P.run(cfg)
        self.assertEqual(tree_hash(self.src), self.before)
        self.assertEqual(rep.stages[0].status, P.OK)

    # ------------------------------------------------------------------ xoá kết quả + Excel
    def test_removable_only_pipeline_dirs_and_excel(self):
        self.assertEqual(P.removable(self.work, self.gather), [])
        os.makedirs(self.work)
        self.assertEqual(P.removable(self.work, self.gather), [])                    # thư mục tình cờ trùng tên: không xoá
        shutil.rmtree(self.work)
        rep = P.run(self.cfg)
        self.assertEqual(P.removable(self.work, self.gather), [self.work, self.gather])
        from pdfread_pro.core import pipeline_excel as PX
        from openpyxl import load_workbook
        path = PX.export(os.path.join(self.tmp, "kq"), rep)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Tổng hợp", "Chi tiết"])
        self.assertEqual(wb["Chi tiết"].max_row, 3 + len(rep.recs))


if __name__ == "__main__":
    unittest.main()
