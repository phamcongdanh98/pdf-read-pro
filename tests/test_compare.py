import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook, load_workbook  # noqa: E402

from pdfread_pro.core import compare as C, report as R, rules1361, scanner  # noqa: E402
from test_report import FIVE, pdf  # noqa: E402

RULES = rules1361.default_rules()


def dossier(root, chibo, folder, codes=FIVE):
    for n in codes:
        pdf(os.path.join(root, chibo, folder, n + ".pdf"))


def make_case(tmp):
    root = os.path.join(tmp, "Hồ sơ số hóa")
    dossier(root, "Chi bo 1 Thai Thong", "056123456789_TranVanThai")            # khớp
    dossier(root, "Chi bo 1 Thai Thong", "056123456780_NguyenVanSai")           # lệch tên (Excel: Nguyễn Văn Sáu)
    dossier(root, "Chi bo 1 Thai Thong", "079111222333_LeThiChuyen")            # khác chi bộ (Excel: chi bộ 2)
    dossier(root, "Chi bo 2 Phuoc Hai", "056999888771_PhamVanGo")               # gõ nhầm (Excel: ...777 1→7?)
    dossier(root, "Chi bo 2 Phuoc Hai", "001222333444_HoVanThua", FIVE[:3])     # ngoài danh sách, thiếu 04 05
    dossier(root, "Chi bo 2 Phuoc Hai", "048555666777_DoThiKhong")              # khớp, CCCD Excel mất số 0
    dossier(root, "Chi bo 2 Phuoc Hai", "056444555666_VoVanKhac")               # CCCD gần giống người khác
    xlsx = os.path.join(tmp, "danh_sach.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "DS dang vien"
    ws.append(["DANH SÁCH ĐẢNG VIÊN ĐÃ SỐ HÓA"])
    ws.append([])
    ws.append(["STT", "Họ và tên", "Ngày sinh", "Số định danh cá nhân", "Tên chi bộ", "Ghi chú"])
    ws.append([1, "Trần Văn Thái", "01/01/1980", "056123456789", "Chi bộ 1 Thái Thông", ""])
    ws.append([2, "Nguyễn Văn Sáu", "", "056123456780", "Chi bộ 1 Thái Thông", ""])
    ws.append([3, "Lê Thị Chuyển", "", "079111222333", "Chi bộ 2 Phước Hải", ""])
    ws.append([4, "Phạm Văn Gõ", "", "056999888777", "Chi bộ 2 Phước Hải", ""])     # thư mục: ...771
    ws.append([5, "Đỗ Thị Không", "", 48555666777, "Chi bộ 2 Phước Hải", ""])       # số, mất số 0 đầu
    ws.append([6, "Bùi Văn Mới", "", "082000111222", "Chi bộ 3 Vĩnh Hải", ""])      # chưa có thư mục + chi bộ mới
    ws.append([7, "Mai Thị Thiếu", "", "056777888999", "Chi bộ 1 Thái Thông", ""])  # chưa có thư mục
    ws.append([8, "Võ Văn Tên Khác", "", "056444555669", "Chi bộ 2 Phước Hải", ""])  # gần giống, tên khác
    ws.append([9, "Trần Văn Thái", "", "056123456789", "Chi bộ 1 Thái Thông", ""])   # trùng trong Excel
    ws.append([10, "", "", "", "", ""])
    ws.append([11, "Dòng lỗi", "", "12345", "Chi bộ 1 Thái Thông", ""])
    wb.save(xlsx)
    return root, xlsx


class CompareTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.root, self.xlsx = make_case(self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_compare(self):
        t = C.read_table(self.xlsx)
        cols = C.guess_columns(t.headers)
        rep = R.build_report(scanner.scan(self.root), RULES)
        return C.compare(t, cols, rep), t, cols

    def by_name(self, cmp, name):
        return [r for r in cmp.rows if r.person and r.person.name == name]

    def test_read_and_guess(self):
        t = C.read_table(self.xlsx)
        self.assertEqual(t.header_row, 3)
        self.assertEqual(t.sheets, ["DS dang vien"])
        self.assertEqual(C.guess_columns(t.headers), {"cccd": 3, "name": 1, "chibo": 4})
        self.assertEqual(C.normalize_cccd(48555666777), ("048555666777", "Đã thêm số 0 đầu (Excel làm mất)"))
        self.assertEqual(C.normalize_cccd("5.6123456789E+10")[0], "056123456789")
        self.assertEqual(C.normalize_cccd("056 123 456 789")[0], "056123456789")
        self.assertEqual(C.normalize_cccd("123456789")[0], "")
        self.assertTrue(C.cccd_near("056999888777", "056999888771"))
        self.assertTrue(C.cccd_near("056123456789", "056123456798"))
        self.assertFalse(C.cccd_near("056123456789", "056123456700"))
        self.assertGreater(C.chibo_similarity("Chi bộ 1 Thái Thông", "Chi bo 1 Thai Thong"), 0.95)
        rows = [(5, ["A", 56123456789, "x"]), (6, ["B", "079111222333", "y"])]
        self.assertEqual(C.guess_columns(["Cột A", "Mã", "Cột C"], rows)["cccd"], 1)      # đoán theo dữ liệu
        self.assertEqual(C.cccd_ratio(rows, 0), 0.0)
        self.assertEqual(C.sample_value(rows, 1), "56123456789")
        with self.assertRaises(C.CompareError):
            C.read_table(os.path.join(self.tmp, "x.xls"))

    def test_categories(self):
        cmp, _, _ = self.run_compare()
        k = lambda name: [r.kind for r in self.by_name(cmp, name)]
        self.assertEqual(k("Nguyễn Văn Sáu"), [C.K_NAME])
        self.assertEqual(k("Lê Thị Chuyển"), [C.K_CHIBO])
        self.assertEqual(k("Phạm Văn Gõ"), [C.K_TYPO])
        self.assertEqual(k("Đỗ Thị Không"), [C.K_MATCH])
        self.assertIn("Đã thêm số 0", self.by_name(cmp, "Đỗ Thị Không")[0].notes[0])
        self.assertEqual(k("Bùi Văn Mới"), [C.K_NO_FOLDER])
        self.assertEqual(k("Mai Thị Thiếu"), [C.K_NO_FOLDER])
        khac = self.by_name(cmp, "Võ Văn Tên Khác")[0]
        self.assertEqual(khac.kind, C.K_NO_FOLDER)
        self.assertTrue(any("gần giống" in n and "họ tên khác" in n for n in khac.notes))
        thai = self.by_name(cmp, "Trần Văn Thái")
        self.assertEqual(sorted(r.kind for r in thai), sorted([C.K_MATCH, C.K_NO_FOLDER]))   # dòng trùng → không có thư mục
        self.assertTrue(all(any("Trùng CCCD trong Excel" in n for n in r.notes) for r in thai))
        out = [r for r in cmp.rows if r.kind == C.K_NOT_IN_LIST]
        self.assertEqual(sorted(r.dossier.folder for r in out), ["001222333444_HoVanThua", "056444555666_VoVanKhac"])
        self.assertEqual([s[0] for s in cmp.skipped], [])                          # dòng 'Dòng lỗi' có tên → giữ, ghi chú
        loi = self.by_name(cmp, "Dòng lỗi")[0]
        self.assertIn("không hợp lệ", loi.notes[0])
        self.assertEqual(cmp.chibo_map["Chi bộ 3 Vĩnh Hải"], "")
        self.assertEqual(cmp.chibo_map["Chi bộ 2 Phước Hải"], "Chi bo 2 Phuoc Hai")
        g = {x["chibo"]: x for x in C.summary_by_chibo(cmp)}
        self.assertEqual((g["Chi bộ 1 Thái Thông"]["excel"], g["Chi bộ 1 Thái Thông"]["match"]), (5, 1))

    def test_actions_and_undo(self):
        cmp, t, cols = self.run_compare()
        before = sorted(os.listdir(os.path.join(self.root, "Chi bo 1 Thai Thong")))
        renames = C.plan_renames(cmp)
        self.assertIn("056123456780_NguyenVanSau", [n for _, n in renames])
        self.assertIn("056999888777_PhamVanGo", [n for _, n in renames])           # sửa cả số CCCD gõ nhầm
        ops = C.apply_ops(renames=renames)
        self.assertTrue(all(o.ok for o in ops), [o.message for o in ops])
        cmp2 = C.compare(t, cols, R.build_report(scanner.scan(self.root), RULES))
        moves = C.plan_moves(cmp2)
        self.assertEqual([os.path.basename(os.path.dirname(d)) for _, d in moves], ["Chi bo 2 Phuoc Hai"])
        ops += C.apply_ops(moves=moves)
        creates = C.plan_creates(cmp2)
        self.assertIn(os.path.join(self.root, "Chi bộ 3 Vĩnh Hải", "082000111222_BuiVanMoi"), creates)
        ops += C.apply_ops(creates=creates)
        cmp3 = C.compare(t, cols, R.build_report(scanner.scan(self.root), RULES))
        self.assertEqual(cmp3.count(C.K_NAME) + cmp3.count(C.K_CHIBO) + cmp3.count(C.K_TYPO), 0)
        C.undo_ops(ops)
        self.assertEqual(sorted(os.listdir(os.path.join(self.root, "Chi bo 1 Thai Thong"))), before)
        self.assertFalse(os.path.exists(os.path.join(self.root, "Chi bộ 3 Vĩnh Hải")))

    def test_export(self):
        from pdfread_pro.core.compare_excel import export_compare
        cmp, _, _ = self.run_compare()
        out = export_compare(cmp, os.path.join(self.tmp, "doi_chieu"))
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Tổng hợp chi bộ", "Đối chiếu chi tiết", "Chưa có thư mục",
                                         "Ngoài danh sách", "Cần sửa", "Giải thích"])
        ws = wb["Đối chiếu chi tiết"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5) if r[0].value]
        self.assertEqual(len(rows), len(cmp.rows))
        sau = next(r for r in rows if r[3] == "Nguyễn Văn Sáu")
        self.assertEqual((sau[2], sau[7]), ("056123456780", C.K_NAME))
        self.assertIn("Đổi tên thư mục → 056123456780_NguyenVanSau", sau[15])
        miss = wb["Chưa có thư mục"]
        self.assertEqual(miss.max_row - 4, 5)    # Bùi Văn Mới, Mai Thị Thiếu, Võ Văn Tên Khác, dòng trùng, Dòng lỗi


if __name__ == "__main__":
    unittest.main()
