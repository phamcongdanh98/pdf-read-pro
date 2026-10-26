import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pymupdf as fitz  # noqa: E402

from pdfread_pro.core import report as R, rules1361, scanner  # noqa: E402

RULES = rules1361.default_rules()


def pdf(path, pages=1):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = fitz.open()
    for _ in range(pages):
        d.new_page()
    d.save(path)


FIVE = ["01.Ly_lich_nguoi_xin_vao_Dang", "02.Ly_lich_dang_vien", "03.Phieu_dang_vien",
        "04.Phieu_bo_sung_ho_so_dang_vien", "05.QD_ket_nap_dang_vien"]


def make_tree(root):
    cb1 = os.path.join(root, "Chi bo 1 Phuoc An Hoa")
    cb2 = os.path.join(root, "Chi bo Phu Tho Cuong")
    ok = os.path.join(cb1, "056123456789_PhamVanA")                 # đạt chuẩn
    for n in FIVE:
        pdf(os.path.join(ok, n + ".pdf"), 2)
    miss = os.path.join(cb1, "056987654321_NguyenThiB")             # thiếu 03, 05; 1 file sai tên
    pdf(os.path.join(miss, FIVE[0] + ".signed.pdf"))
    pdf(os.path.join(miss, "2.ly lich.pdf"))
    pdf(os.path.join(miss, FIVE[3] + ".pdf"))
    bad = os.path.join(cb2, "079111222333 - Trần văn c")            # sai tên thư mục, có file thừa
    for n in FIVE:
        pdf(os.path.join(bad, n + ".pdf"))
    pdf(os.path.join(bad, "06.QD_ket_nap_lai.pdf"))
    pdf(os.path.join(bad, "scan moi.pdf"))
    os.makedirs(os.path.join(cb2, "045000111222_LeThiD"))           # trống
    dup = os.path.join(cb2, "056123456789_PhamVanA")                # trùng CCCD với chi bộ 1
    for n in FIVE:
        pdf(os.path.join(dup, n + ".pdf"))
    pdf(os.path.join(cb2, "lac_cho.pdf"))                           # file nằm thẳng trong chi bộ
    return cb1, cb2


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.root = os.path.join(self.tmp, "Hồ sơ số hóa")
        make_tree(self.root)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def report(self, root=None):
        return R.build_report(scanner.scan(root or self.root), RULES)

    def find(self, rep, chibo, folder):
        return next(d for d in rep.dossiers if d.chibo == chibo and d.folder == folder)

    def test_classification(self):
        rep = self.report()
        self.assertEqual(rep.mode, "root")
        self.assertEqual(rep.chibos, ["Chi bo 1 Phuoc An Hoa", "Chi bo Phu Tho Cuong"])
        ok = self.find(rep, "Chi bo 1 Phuoc An Hoa", "056123456789_PhamVanA")
        # trùng CCCD với hồ sơ ở chi bộ 2 → cần xử lý
        self.assertEqual(ok.status, R.ST_FIX)
        self.assertEqual((ok.cccd, ok.name, ok.pages), ("056123456789", "Pham Van A", 10))
        miss = self.find(rep, "Chi bo 1 Phuoc An Hoa", "056987654321_NguyenThiB")
        self.assertEqual(miss.status, R.ST_MISSING)
        self.assertEqual(miss.missing, ["03", "05"])
        self.assertTrue(any("Sai tên file: 2.ly lich.pdf" in i.text for i in miss.issues))
        bad = self.find(rep, "Chi bo Phu Tho Cuong", "079111222333 - Trần văn c")
        self.assertEqual(bad.status, R.ST_NAME)
        self.assertEqual(bad.suggestion, "079111222333_TranVanC")
        self.assertTrue(any("Sai tên file (giấy tờ không bắt buộc): 06." in i.text for i in bad.issues))
        self.assertTrue(any("Không xác định loại" in i.text for i in bad.issues))
        empty = self.find(rep, "Chi bo Phu Tho Cuong", "045000111222_LeThiD")
        self.assertEqual(empty.status, R.ST_EMPTY)
        self.assertEqual([s.rel.split(os.sep)[-1] for s in rep.stray], ["lac_cho.pdf"])

    def test_scan_one_chibo_and_one_dossier(self):
        rep = self.report(os.path.join(self.root, "Chi bo 1 Phuoc An Hoa"))
        self.assertEqual(rep.mode, "chibo")
        self.assertEqual(len(rep.dossiers), 2)
        self.assertEqual(self.find(rep, "Chi bo 1 Phuoc An Hoa", "056123456789_PhamVanA").status, R.ST_OK)
        rep = self.report(os.path.join(self.root, "Chi bo 1 Phuoc An Hoa", "056123456789_PhamVanA"))
        self.assertEqual(rep.mode, "dossier")
        self.assertEqual([d.status for d in rep.dossiers], [R.ST_OK])

    def test_required_codes_configurable(self):
        rep = R.build_report(scanner.scan(self.root), RULES, R.parse_required("01, 04"))
        miss = self.find(rep, "Chi bo 1 Phuoc An Hoa", "056987654321_NguyenThiB")
        self.assertEqual(miss.missing, [])
        self.assertEqual(miss.status, R.ST_WARN)          # file 02 nằm ngoài danh mục → chỉ cảnh báo
        with self.assertRaises(ValueError):
            R.parse_required("abc")

    def test_folder_names(self):
        self.assertEqual(R.suggest_folder_name("056123456789 Phạm Văn A"), "056123456789_PhamVanA")
        self.assertEqual(R.suggest_folder_name("056123456789_pham_van_a"), "056123456789_PhamVanA")
        ok, *_ = R.check_folder("056123456789_PHAMVANA")
        self.assertFalse(ok)
        ok, _, _, issues, sug = R.check_folder("05612345678_PhamVanA")
        self.assertFalse(ok)
        self.assertEqual(sug, "")
        self.assertIn("11 chữ số", issues[0].text)
        ok, _, _, issues, _ = R.check_folder("999123456789_PhamVanA")
        self.assertTrue(ok)
        self.assertEqual(issues[0].level, R.LV_WARN)

    def test_excel_export(self):
        from openpyxl import load_workbook
        from pdfread_pro.core.report_excel import export_excel
        rep = self.report()
        out = export_excel(rep, os.path.join(self.tmp, "bao_cao"))
        self.assertTrue(out.endswith(".xlsx") and os.path.exists(out))
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Tổng hợp chi bộ", "Danh sách toàn bộ", "Danh sách thiếu", "Giải thích"])
        # tổng hợp: chi bộ 1 có 2 hồ sơ, đủ 1, thiếu 1 (thiếu 03 và 05)
        t = wb["Tổng hợp chi bộ"]
        self.assertEqual([t.cell(row=6, column=c).value for c in range(2, 11)],
                         ["Chi bo 1 Phuoc An Hoa", 2, 1, 1, 0, 0, 1, 0, 1])
        # danh sách toàn bộ: đủ 5 người (cả đủ lẫn thiếu)
        ws = wb["Danh sách toàn bộ"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5, max_row=9)]
        self.assertEqual(len([r for r in rows if r[0]]), 5)
        b = next(r for r in rows if r[2] == "056987654321")
        self.assertEqual(b[3], "Nguyen Thi B")
        self.assertEqual(b[4:9], ["Có", "Có", "Thiếu", "Có", "Thiếu"])
        self.assertEqual(b[12:14], ["Thiếu", "03 – Phiếu đảng viên; 05 – Quyết định kết nạp đảng viên"])
        self.assertIn("Sai tên file: 2.ly lich.pdf", b[14])
        a = next(r for r in rows if r[2] == "056123456789" and r[1] == "Chi bo 1 Phuoc An Hoa")
        self.assertEqual((a[12], a[13]), ("Đủ", None))
        self.assertIn("Trùng số CCCD", a[14])
        # danh sách thiếu: chỉ người thiếu (Nguyen Thi B + hồ sơ trống Le Thi D)
        th = wb["Danh sách thiếu"]
        names = [th.cell(row=r, column=4).value for r in range(5, 8)]
        self.assertEqual(sorted(n for n in names if n), ["Le Thi D", "Nguyen Thi B"])
        # kèm báo cáo nâng cao
        wb2 = load_workbook(export_excel(rep, os.path.join(self.tmp, "bao_cao_nc.xlsx"), advanced=True))
        self.assertEqual(wb2.sheetnames[4:], ["Chi tiết file", "Thống kê giấy tờ", "Vấn đề cần xử lý"])
        tk = {r[0].value: r[4].value for r in wb2["Thống kê giấy tờ"].iter_rows(min_row=5)}
        self.assertEqual((tk["01"], tk["03"], len(tk)), (4, 3, len(RULES)))

    def test_signed_filter(self):
        cb = os.path.join(self.tmp, "HS", "Chi bo Ky so")
        pair = os.path.join(cb, "056111222333_NguyenVanKy")          # đủ 5 bản chưa ký + đủ 5 bản đã ký
        for n in FIVE:
            pdf(os.path.join(pair, n + ".pdf"))
            pdf(os.path.join(pair, n + ".signed.pdf"))
        half = os.path.join(cb, "056111222444_TranThiChua")         # 01 mới có bản chưa ký
        pdf(os.path.join(half, FIVE[0] + ".pdf"))
        for n in FIVE[1:]:
            pdf(os.path.join(half, n + ".signed.pdf"))
        only = os.path.join(cb, "056111222555_LeVanDa")             # chỉ có bản đã ký
        for n in FIVE:
            pdf(os.path.join(only, n + ".signed.pdf"))
        res = scanner.scan(os.path.join(self.tmp, "HS"))

        def get(sf):
            rep = R.build_report(res, RULES, sign_filter=sf)
            return rep, {d.folder.split("_")[1]: d for d in rep.dossiers}

        rep, d = get(R.SIGN_ALL)
        self.assertTrue(all(x.complete for x in d.values()))
        self.assertEqual(d["NguyenVanKy"].notes, [])                  # cặp chưa ký/đã ký không bị báo trùng
        rep, d = get(R.SIGN_ONLY)
        self.assertTrue(d["NguyenVanKy"].complete and d["LeVanDa"].complete)
        self.assertEqual(d["TranThiChua"].missing, ["01"])
        self.assertEqual(R.missing_hint(d["TranThiChua"], "01", rep.sign_filter), " (mới có bản chưa ký)")
        rep, d = get(R.SIGN_NONE)
        self.assertTrue(d["NguyenVanKy"].complete)
        self.assertEqual(d["TranThiChua"].missing, ["02", "03", "04", "05"])
        self.assertEqual(d["LeVanDa"].status, R.ST_EMPTY)
        self.assertEqual(d["LeVanDa"].notes, [])
        from openpyxl import load_workbook
        from pdfread_pro.core.report_excel import export_excel
        rep, _ = get(R.SIGN_ONLY)
        wb = load_workbook(export_excel(rep, os.path.join(self.tmp, "ky.xlsx")))
        self.assertIn("Chỉ file đã ký", wb["Danh sách toàn bộ"]["A2"].value)
        row = next(r for r in wb["Danh sách thiếu"].iter_rows(min_row=5, values_only=True) if r[0])
        self.assertEqual(row[13], "01 – Lý lịch người xin vào Đảng (mới có bản chưa ký)")
        with self.assertRaises(ValueError):
            R.build_report(res, RULES, sign_filter="abc")

    def test_rename_folders_and_undo(self):
        rep = self.report()
        bad = self.find(rep, "Chi bo Phu Tho Cuong", "079111222333 - Trần văn c")
        res = R.rename_folders([(bad.path, bad.suggestion)])
        self.assertTrue(res[0].ok, res[0].message)
        self.assertTrue(os.path.isdir(os.path.join(os.path.dirname(bad.path), "079111222333_TranVanC")))
        rep2 = self.report()
        self.assertEqual(self.find(rep2, "Chi bo Phu Tho Cuong", "079111222333_TranVanC").status, R.ST_WARN)
        R.undo_renames(res)
        self.assertTrue(os.path.isdir(bad.path))
        # không ghi đè thư mục đã tồn tại
        cb2 = os.path.dirname(bad.path)
        res = R.rename_folders([(bad.path, "045000111222_LeThiD")])
        self.assertFalse(res[0].ok)
        # đổi chỉ khác hoa/thường
        p = os.path.join(cb2, "045000111222_LeThiD")
        res = R.rename_folders([(p, "045000111222_LETHID")])
        self.assertTrue(res[0].ok, res[0].message)
        self.assertIn("045000111222_LETHID", os.listdir(cb2))


if __name__ == "__main__":
    unittest.main()
