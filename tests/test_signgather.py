import os
import shutil
import sys
import tempfile
import time
import unittest
from unittest import mock

import fitz

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pdfread_pro.core import scanner, signgather as G, signgather_excel as GX  # noqa: E402
from pdfread_pro.core.models import FileInfo  # noqa: E402
from make_signed_sample import sign_pdf  # noqa: E402

A, B = "056123456789_TranVanThai", "079111222333_LeThiChuyen"
A01 = f"01.Ly_lich_nguoi_xin_vao_Dang_{A}"
B01 = f"01.Ly_lich_nguoi_xin_vao_Dang_{B}"


def touch(path, data=b"%PDF-1.4 x"):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(data)


def make_pdf(path, pages=1, text="x"):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = fitz.open()
    for i in range(pages):
        d.new_page().insert_text((72, 72), f"{text} {i}")
    d.save(path)
    d.close()


class SignGatherTests(unittest.TestCase):
    def setUp(self):
        G.BUSY_SECONDS = 0
        G._CACHE.clear()
        self.tmp = tempfile.mkdtemp()
        self.root = os.path.join(self.tmp, "Chi bo 1")
        self.a = os.path.join(self.root, A)
        self.b = os.path.join(self.root, B)
        touch(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.pdf"), b"%PDF a01")
        touch(os.path.join(self.a, "02.Ly_lich_dang_vien.pdf"), b"%PDF a02")
        touch(os.path.join(self.a, "02.Ly_lich_dang_vien.signed.pdf"), b"%PDF a02 signed")    # đã ký → bỏ qua
        touch(os.path.join(self.b, "01.Ly_lich_nguoi_xin_vao_Dang.pdf"), b"%PDF b01")
        touch(os.path.join(self.root, "Khac", "ghi_chu.pdf"))                                  # không phải hồ sơ
        self.gdir = os.path.join(self.tmp, "Gom")

    def tearDown(self):
        G.BUSY_SECONDS = 5
        shutil.rmtree(self.tmp, ignore_errors=True)

    def files(self, root=None):
        return scanner.scan(root or self.root).files

    def gather_all(self, **kw):
        return G.gather(G.plan_gather(self.files(), self.gdir, **kw), self.gdir)

    # ---------------------------------------------------------------- luồng cơ bản (không kiểm tra chữ ký thật)
    def test_full_cycle(self):
        plan = G.plan_gather(self.files(), self.gdir)
        st = {os.path.basename(it.src) + "@" + it.folder: it.status for it in plan}
        self.assertEqual(st[f"02.Ly_lich_dang_vien.pdf@{A}"], G.G_SIGNED)
        self.assertEqual(st["ghi_chu.pdf@Khac"], G.G_NOT_DOSSIER)
        new = sorted(it.new_name for it in plan if it.status == G.G_NEW)
        self.assertEqual(new, [A01 + ".pdf", B01 + ".pdf"])
        j = G.gather(plan, self.gdir)
        self.assertEqual(len(j.copied), 2)
        self.assertTrue(os.path.isfile(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.pdf")))   # gốc còn nguyên
        self.assertFalse([n for n in os.listdir(self.gdir) if n.endswith(".part")])                  # không file dở
        self.assertTrue(all(it.status != G.G_NEW for it in G.plan_gather(self.files(), self.gdir)
                            if it.status != G.G_NOT_DOSSIER and it.status != G.G_SIGNED))

        touch(os.path.join(self.gdir, A01 + ".signed.pdf"), b"%PDF a01 signed")
        ret = {os.path.basename(it.path): it for it in G.plan_return(self.gdir, [self.root], verify=False)}
        s = ret[A01 + ".signed.pdf"]
        self.assertEqual((s.status, s.dest), (G.R_SIGNED, os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf")))
        self.assertEqual(ret[A01 + ".pdf"].status, G.R_DUP)
        self.assertEqual(ret[B01 + ".pdf"].status, G.R_WAIT)
        j2 = G.do_return(list(ret.values()))
        self.assertTrue(os.path.isfile(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf")))
        left = sorted(n for n in os.listdir(self.gdir) if n.endswith(".pdf"))
        self.assertEqual(left, [B01 + ".pdf"])                                                   # B chờ ký
        G.undo(j2)
        self.assertFalse(os.path.exists(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf")))
        self.assertEqual(len([n for n in os.listdir(self.gdir) if n.endswith(".pdf")]), 3)

    def test_no_manifest_and_conflict(self):
        """Không có sổ ghi (vd chép thư mục gom sang máy khác): nhận hồ sơ theo đuôi _CCCD_HoVaTen; không ghi đè."""
        os.makedirs(self.gdir)
        touch(os.path.join(self.gdir, B01 + "_signed.pdf"), b"new")
        touch(os.path.join(self.b, "01.Ly_lich_nguoi_xin_vao_Dang_signed.pdf"), b"old khac")
        touch(os.path.join(self.gdir, "lung_tung.pdf"))
        ret = {os.path.basename(it.path): it for it in G.plan_return(self.gdir, [self.root], verify=False)}
        self.assertEqual(ret[B01 + "_signed.pdf"].status, G.R_CONFLICT)
        self.assertEqual(ret["lung_tung.pdf"].status, G.R_UNKNOWN)
        G.do_return(list(ret.values()))
        with open(os.path.join(self.b, "01.Ly_lich_nguoi_xin_vao_Dang_signed.pdf"), "rb") as fh:
            self.assertEqual(fh.read(), b"old khac")                                      # không bị ghi đè

    def test_copy_out_full_set(self):
        """Chép ra thư mục mới: giữ cây Chi bộ/CCCD_HoVaTen, hồ sơ gốc + thư mục gom không đổi, kèm file khác."""
        self.gather_all()
        touch(os.path.join(self.gdir, A01 + ".signed.pdf"), b"%PDF a01 k")
        out = os.path.join(self.tmp, "Da ky")
        before_src, before_g = sorted(os.listdir(self.a)), sorted(os.listdir(self.gdir))
        plan = G.plan_return(self.gdir, [self.tmp], out, full=True, verify=False)
        st = {(it.folder, os.path.basename(it.dest)): it.status for it in plan if it.dest}
        oa = os.path.join(out, "Chi bo 1", A)
        self.assertEqual(st[(A, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf")], G.R_SIGNED)
        self.assertEqual(st[(A, "01.Ly_lich_nguoi_xin_vao_Dang.pdf")], G.R_COPY)
        self.assertEqual(st[(A, "02.Ly_lich_dang_vien.signed.pdf")], G.R_EXTRA)
        self.assertIn(G.R_WAIT, [it.status for it in plan if it.folder == B])
        j = G.do_return(plan, copy=True)
        self.assertEqual(sorted(os.listdir(oa)), ["01.Ly_lich_nguoi_xin_vao_Dang.pdf",
                                                  "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf",
                                                  "02.Ly_lich_dang_vien.pdf", "02.Ly_lich_dang_vien.signed.pdf"])
        self.assertFalse(os.path.exists(os.path.join(out, "Chi bo 1", B)))
        self.assertEqual(sorted(os.listdir(self.a)), before_src)
        self.assertEqual(sorted(os.listdir(self.gdir)), before_g)
        again = G.plan_return(self.gdir, [self.tmp], out, full=True, verify=False)
        self.assertFalse([it for it in again if it.status in G.COPY_ACTIONS])
        self.assertEqual(G.undo(j), 4)
        self.assertEqual(os.listdir(oa), [])

    def test_manifest_hidden_broken_or_missing(self):
        """Sổ _gom_ky_so.json: ẩn (kể cả sau lần ghi thứ 2); hỏng / bị xoá vẫn trả về được theo đuôi tên."""
        self.gather_all()
        touch(os.path.join(self.a, "02b.Phu_luc.pdf"), b"%PDF a02b")
        self.gather_all()
        man = os.path.join(self.gdir, G.MANIFEST)
        if os.name == "nt":
            import ctypes
            self.assertTrue(ctypes.windll.kernel32.GetFileAttributesW(man) & 0x2)
        touch(os.path.join(self.gdir, A01 + ".signed.pdf"), b"k")
        for broken in (b"{khong phai json", None):
            if os.path.exists(man):
                os.remove(man)                    # file ẩn không mở "wb" trực tiếp được trên Windows
            if broken is not None:
                with open(man, "wb") as fh:
                    fh.write(broken)
            st = {os.path.basename(it.path): it.status for it in G.plan_return(self.gdir, [self.root], verify=False)}
            self.assertEqual(st[A01 + ".signed.pdf"], G.R_SIGNED)
            self.assertNotIn(G.MANIFEST, st)

    def test_guess_root_without_scan(self):
        G.gather(G.plan_gather(scanner.scan(self.tmp).files, self.gdir), self.gdir)
        self.assertEqual(G.guess_root(self.gdir), self.root)
        touch(os.path.join(self.gdir, A01 + ".signed.pdf"), b"k")
        out = os.path.join(self.tmp, "Da ky")
        plan = G.plan_return(self.gdir, [G.guess_root(self.gdir)], out, full=False, verify=False)
        signed = [it for it in plan if it.status == G.R_SIGNED]
        self.assertEqual(signed[0].dest, os.path.join(out, A, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf"))
        self.assertEqual(G.guess_root(os.path.join(self.tmp, "khong_co")), "")
        G._save_manifest(self.gdir, {"items": {A01 + ".pdf": "hong"}})
        st = {os.path.basename(it.path): it.status for it in G.plan_return(self.gdir, [self.root], verify=False)}
        self.assertEqual(st[A01 + ".signed.pdf"], G.R_SIGNED)

    def test_plan_return_fast_on_thousands(self):
        os.makedirs(self.gdir)
        for i in range(1500):
            for tail in ("", ".signed"):
                touch(os.path.join(self.gdir, f"{i:04d}.Giay_to_{A}{tail}.pdf"), b"x")
        t = time.perf_counter()
        plan = G.plan_return(self.gdir, [self.root], verify=False)
        self.assertLess(time.perf_counter() - t, 3.0)
        self.assertEqual(sum(1 for it in plan if it.status == G.R_SIGNED), 1500)

    def test_undo_gather(self):
        j = self.gather_all()
        self.assertEqual(G.undo(j), 2)
        self.assertEqual([n for n in os.listdir(self.gdir) if n.endswith(".pdf")], [])

    # ---------------------------------------------------------------- A1 / A2: chữ ký thật + số trang
    def _real(self):
        """Hồ sơ PDF thật: A01 3 trang, B01 2 trang, A03 1 trang → gom."""
        make_pdf(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.pdf"), 3, "A")
        make_pdf(os.path.join(self.b, "01.Ly_lich_nguoi_xin_vao_Dang.pdf"), 2, "B")
        make_pdf(os.path.join(self.a, "03.Phieu_dang_vien.pdf"), 1, "A3")
        self.gather_all()

    def test_verify_signature_and_pages(self):
        self._real()
        g = lambda n: os.path.join(self.gdir, n)  # noqa: E731
        sign_pdf(g(A01 + ".pdf"), g(A01 + ".signed.pdf"), "TRAN VAN THAI")              # ký thật → nhận
        shutil.copy(g(B01 + ".pdf"), g(B01 + ".signed.pdf"))                            # chỉ đổi tên → KHÔNG nhận
        a03 = f"03.Phieu_dang_vien_{A}"
        make_pdf(g("tmp3.pdf"), 4, "khac")                                              # ký nhầm file 4 trang
        sign_pdf(g("tmp3.pdf"), g(a03 + ".signed.pdf"))
        os.remove(g("tmp3.pdf"))
        plan = G.plan_return(self.gdir, [self.root])
        st = {os.path.basename(it.path): it for it in plan}
        self.assertEqual(st[A01 + ".signed.pdf"].status, G.R_SIGNED)
        self.assertIn("TRAN VAN THAI", st[A01 + ".signed.pdf"].note)
        self.assertEqual(st[A01 + ".pdf"].status, G.R_DUP)
        self.assertEqual(st[B01 + ".signed.pdf"].status, G.R_NOSIG)
        self.assertEqual(st[B01 + ".pdf"].status, G.R_WAIT)                             # bản sao giữ lại, không xoá
        self.assertEqual(st[a03 + ".signed.pdf"].status, G.R_PAGES)
        self.assertIn("4 trang", st[a03 + ".signed.pdf"].note)
        # chữ ký hỏng: sửa 1 byte trong vùng đã ký (ngày /M) – PDF vẫn mở được
        with open(g(A01 + ".signed.pdf"), "rb") as fh:
            data = fh.read()
        with open(g(A01 + ".signed.pdf"), "wb") as fh:
            fh.write(data.replace(b"D:2026", b"D:2025", 1))
        st = {os.path.basename(it.path): it.status for it in G.plan_return(self.gdir, [self.root])}
        self.assertEqual(st[A01 + ".signed.pdf"], G.R_BADSIG)
        self.assertEqual(st[A01 + ".pdf"], G.R_WAIT)
        # file rỗng
        touch(g(B01 + ".signed.pdf"), b"")
        st = {os.path.basename(it.path): it.status for it in G.plan_return(self.gdir, [self.root])}
        self.assertEqual(st[B01 + ".signed.pdf"], G.R_BAD)

    def test_inplace_subfolder_and_other_tails(self):
        """B6: phần mềm ký đè lên chính file / ghi ra Signed/ cùng tên / đuôi _ky."""
        self._real()
        g = lambda *n: os.path.join(self.gdir, *n)  # noqa: E731
        tmp = g("x.tmp")
        sign_pdf(g(A01 + ".pdf"), tmp)
        os.replace(tmp, g(A01 + ".pdf"))                                               # ký đè
        os.makedirs(g("Signed"))
        sign_pdf(g(B01 + ".pdf"), g("Signed", B01 + ".pdf"))                            # ghi ra Signed/ cùng tên
        a03 = f"03.Phieu_dang_vien_{A}"
        sign_pdf(g(a03 + ".pdf"), g(a03 + "_ky.pdf"))                                   # đuôi _ky
        plan = G.plan_return(self.gdir, [self.root])
        dest = {os.path.basename(it.path) + ("@S" if "Signed" in it.path else ""): it for it in plan}
        self.assertEqual(dest[A01 + ".pdf"].status, G.R_SIGNED)
        self.assertTrue(dest[A01 + ".pdf"].dest.endswith("01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf"))
        self.assertEqual(dest[B01 + ".pdf@S"].status, G.R_SIGNED)
        self.assertEqual(dest[B01 + ".pdf"].status, G.R_DUP)
        self.assertEqual(dest[a03 + "_ky.pdf"].status, G.R_SIGNED)
        self.assertTrue(dest[a03 + "_ky.pdf"].dest.endswith("03.Phieu_dang_vien.signed.pdf"))
        G.do_return(plan)
        self.assertTrue(os.path.isfile(os.path.join(self.b, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf")))
        self.assertTrue(os.path.isfile(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.signed.pdf")))

    # ---------------------------------------------------------------- A3 / A6 / A7
    def test_busy_and_changed_original(self):
        self._real()
        sign_pdf(os.path.join(self.gdir, A01 + ".pdf"), os.path.join(self.gdir, A01 + ".signed.pdf"))
        G.BUSY_SECONDS = 60                                                             # vừa ghi xong → đang ký
        st = {os.path.basename(it.path): it.status for it in G.plan_return(self.gdir, [self.root])}
        self.assertEqual(st[A01 + ".signed.pdf"], G.R_BUSY)
        self.assertEqual(st[A01 + ".pdf"], G.R_BUSY)
        G.BUSY_SECONDS = 0
        make_pdf(os.path.join(self.a, "01.Ly_lich_nguoi_xin_vao_Dang.pdf"), 3, "SCAN LAI")   # bản gốc bị thay
        st = {os.path.basename(it.path): it.status for it in G.plan_return(self.gdir, [self.root])}
        self.assertEqual(st[A01 + ".signed.pdf"], G.R_CHANGED)
        # tách ra _Bản ký lỗi → gom lại được bản mới
        pairs = G.plan_regroup(G.plan_return(self.gdir, [self.root]), self.gdir)
        G.regroup(pairs)
        again = [it for it in G.plan_gather(self.files(), self.gdir) if it.new_name == A01 + ".pdf"]
        self.assertEqual(again[0].status, G.G_NEW)

    def test_duplicate_dossier_name(self):
        other = os.path.join(self.tmp, "Chi bo 2", A)
        touch(os.path.join(other, "01.Ly_lich_nguoi_xin_vao_Dang.pdf"), b"%PDF khac")
        plan = G.plan_gather(scanner.scan(self.tmp).files, self.gdir)
        dup = [it for it in plan if it.folder == A and it.status == G.G_DUPNAME]
        self.assertEqual(len(dup), 3)                     # 2 file chưa ký ở Chi bo 1 + 1 file ở Chi bo 2
        self.assertIn("Chi bo", dup[0].note)
        self.assertFalse([it for it in plan if it.folder == A and it.status == G.G_NEW])

    # ---------------------------------------------------------------- A4 / A8
    def test_long_path_and_part_file(self):
        deep = os.path.join(self.tmp, *(["Thu_muc_rat_dai_" + "x" * 30] * 6), A)
        src = os.path.join(deep, "01.Ly_lich_nguoi_xin_vao_Dang.pdf")
        self.assertGreater(len(src), 270)
        os.makedirs(G._lp(deep))
        with open(G._lp(src), "wb") as fh:
            fh.write(b"%PDF long")
        fi = FileInfo(path=src, rel="", name=os.path.basename(src))
        gdir = os.path.join(os.path.dirname(deep), "Gom")
        plan = G.plan_gather([fi], gdir)
        self.assertIn("đường dẫn dài", plan[0].note)
        G.gather(plan, gdir)
        with open(G._lp(os.path.join(gdir, A01 + ".signed.pdf")), "wb") as fh:
            fh.write(b"%PDF long k")
        out = os.path.join(os.path.dirname(deep), "Da ky")
        ret = G.plan_return(gdir, [os.path.dirname(deep)], out, full=False, verify=False)
        j = G.do_return(ret, copy=True)
        self.assertEqual(len(j.copied), 2)
        self.assertTrue(all(os.path.exists(G._lp(p)) for p in j.copied))
        self.assertEqual(G.undo(j), 2)

    def test_not_enough_space(self):
        fake = mock.Mock(free=10)
        with mock.patch.object(G.shutil, "disk_usage", return_value=fake):
            with self.assertRaises(OSError) as cm:
                self.gather_all()
        self.assertIn("không đủ chỗ", str(cm.exception))
        self.assertFalse(os.path.exists(self.gdir))                                    # chưa chép gì

    # ---------------------------------------------------------------- A5 / B2 / B3 / B5
    def test_journal_persist_batch_regroup_excel(self):
        self._real()
        shutil.rmtree(self.gdir)
        plan = G.plan_gather(self.files(), self.gdir, batch=G.BATCH_SIZE, batch_size=2)
        subs = sorted({it.sub for it in plan if it.status == G.G_NEW})
        self.assertEqual(subs, ["Lô 01", "Lô 02"])
        j = G.gather(plan, self.gdir)
        jp = os.path.join(self.tmp, "nhat_ky.json")
        G.save_journal(jp, j)
        j2 = G.load_journal(jp)
        self.assertEqual((j2.kind, j2.copied), ("gather", j.copied))
        self.assertIsNone(G.load_journal(os.path.join(self.tmp, "khong_co.json")))
        # ký 1 file trong lô, file còn lại chờ → tách ra "Ký lại"
        lo = os.path.join(self.gdir, "Lô 01")
        first = sorted(os.listdir(lo))[0]
        sign_pdf(os.path.join(lo, first), os.path.join(lo, first[:-4] + ".signed.pdf"))
        ret = G.plan_return(self.gdir, [self.root])
        self.assertEqual(sum(1 for it in ret if it.status == G.R_SIGNED), 1)
        waits = [it for it in ret if it.status == G.R_WAIT]
        pairs = G.plan_regroup(ret, self.gdir)
        self.assertEqual(len(pairs), len(waits))
        self.assertTrue(all(os.path.basename(os.path.dirname(d)).startswith(G.REDO_DIR) for _, d in pairs))
        jr = G.regroup(pairs)
        self.assertEqual(G.plan_regroup(G.plan_return(self.gdir, [self.root]), self.gdir), [])   # không tách lần 2
        G.undo(jr)
        self.assertTrue(all(os.path.exists(s) for s, _ in pairs))
        # Excel
        rows = [(it.folder, os.path.basename(it.path), it.dest, it.status, G.TEXT[it.status], it.note, it.signed)
                for it in ret]
        xl = GX.export(os.path.join(self.tmp, "kq"), rows, "Kết quả gom ký số", "test")
        from openpyxl import load_workbook
        wb = load_workbook(xl)
        self.assertEqual(wb.sheetnames, ["Theo hồ sơ", "Chi tiết"])
        self.assertEqual(wb["Chi tiết"].max_row, 3 + len(rows))


if __name__ == "__main__":
    unittest.main()
