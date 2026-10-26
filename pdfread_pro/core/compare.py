"""ĐỐI CHIẾU thư mục hồ sơ số hoá với DANH SÁCH HỆ THỐNG (file Excel .xlsx xuất từ hệ thống).

Chỉ cần 3 cột: Số CCCD (khoá chính) · Họ và tên · Chi bộ (không bắt buộc). Tên cột tuỳ file – app tự đoán,
người dùng chỉnh lại được. Kết quả mỗi người: Khớp · Lệch tên · Khác chi bộ · Nghi gõ nhầm CCCD ·
Chưa có thư mục · Ngoài danh sách; kèm tình trạng file 01–05 lấy từ báo cáo hồ sơ.
Thao tác sửa (có hoàn tác): đổi tên thư mục theo Excel · chuyển về đúng chi bộ · tạo thư mục còn thiếu."""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Dict, List, Optional, Sequence, Tuple

from . import report as R
from .renamer import remove_accents

# kết quả đối chiếu
K_MATCH = "Khớp"
K_NAME = "Lệch tên"
K_CHIBO = "Khác chi bộ"
K_TYPO = "Nghi gõ nhầm CCCD"
K_NO_FOLDER = "Chưa có thư mục"
K_NOT_IN_LIST = "Ngoài danh sách"
KIND_ORDER = (K_NO_FOLDER, K_NOT_IN_LIST, K_TYPO, K_NAME, K_CHIBO, K_MATCH)
NEED_FIX = (K_NAME, K_CHIBO, K_TYPO)

_KEYS = {
    "cccd": ("cccd", "can cuoc", "dinh danh", "cmnd", "so cmt", "chung minh"),
    "name": ("ho va ten", "ho ten", "ten dang vien", "hoten", "ho, ten"),
    "chibo": ("chi bo", "don vi", "to chuc dang", "chibo"),
}


class CompareError(Exception):
    pass


# ------------------------------------------------------------------ đọc Excel

@dataclass
class Table:
    path: str
    sheet: str
    sheets: List[str]
    header_row: int                    # 1-based
    headers: List[str]
    rows: List[Tuple[int, list]]       # (số dòng Excel, giá trị)


def _norm(s) -> str:
    return re.sub(r"\s+", " ", remove_accents(str(s or "")).lower()).strip()


def _header_score(values: list) -> int:
    score = 0
    for v in values:
        t = _norm(v)
        if not t or len(t) > 60:
            continue
        score += sum(1 for keys in _KEYS.values() if any(k in t for k in keys))
    return score


def read_table(path: str, sheet: Optional[str] = None, header_row: Optional[int] = None) -> Table:
    """Đọc 1 sheet. header_row = None → tự dò dòng tiêu đề trong 30 dòng đầu."""
    if not path.lower().endswith((".xlsx", ".xlsm")):
        raise CompareError("Chỉ hỗ trợ file Excel .xlsx (Excel 2007 trở lên).")
    if not os.path.isfile(path):
        raise CompareError("Không tìm thấy file Excel.")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise CompareError("File đang mở ở chương trình khác – hãy đóng file rồi thử lại.")
    except Exception as e:
        raise CompareError(f"Không đọc được file Excel: {e}")
    try:
        sheets = wb.sheetnames
        ws = wb[sheet] if sheet in sheets else wb[sheets[0]]
        all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not all_rows:
        raise CompareError("Sheet trống.")
    if header_row is None:
        best, header_row = 0, 1
        for i, r in enumerate(all_rows[:30], 1):
            sc = _header_score(r)
            if sc > best:
                best, header_row = sc, i
    header_row = max(1, min(header_row, len(all_rows)))
    raw = all_rows[header_row - 1]
    width = max((len(r) for r in all_rows[header_row - 1:header_row + 200]), default=len(raw))
    headers = []
    for i in range(width):
        v = raw[i] if i < len(raw) else None
        t = re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""
        headers.append(t or f"Cột {_col_letter(i)}")
    rows = []
    for n, r in enumerate(all_rows[header_row:], header_row + 1):
        vals = list(r) + [None] * (width - len(r))
        if any(v not in (None, "") for v in vals):
            rows.append((n, vals[:width]))
    return Table(path, ws.title, sheets, header_row, headers, rows)


def _col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def guess_columns(headers: Sequence[str], rows: Optional[Sequence[Tuple[int, list]]] = None
                  ) -> Dict[str, Optional[int]]:
    """Đoán cột theo tên tiêu đề; có rows thì cột CCCD được chọn lại theo dữ liệu (cột có nhiều số CCCD nhất)
    khi tên tiêu đề không rõ hoặc cột đoán theo tên không chứa CCCD."""
    out: Dict[str, Optional[int]] = {"cccd": None, "name": None, "chibo": None}
    for key, words in _KEYS.items():
        for i, h in enumerate(headers):
            t = _norm(h)
            if any(w in t for w in words) and i not in out.values():
                out[key] = i
                break
    if rows:
        cur = out["cccd"]
        if cur is None or cccd_ratio(rows, cur) < 0.5:
            best = max(range(len(headers)), key=lambda i: cccd_ratio(rows, i), default=None)
            if best is not None and cccd_ratio(rows, best) >= 0.5 and best not in (out["name"], out["chibo"]):
                out["cccd"] = best
    return out


def cccd_ratio(rows: Sequence[Tuple[int, list]], col: Optional[int], sample: int = 300) -> float:
    """Tỉ lệ ô có dữ liệu trong cột là số CCCD hợp lệ (để báo người dùng chọn nhầm cột)."""
    if col is None:
        return 0.0
    vals = [v[col] for _, v in rows[:sample] if col < len(v) and v[col] not in (None, "")]
    if not vals:
        return 0.0
    return sum(1 for x in vals if normalize_cccd(x)[0]) / len(vals)


def sample_value(rows: Sequence[Tuple[int, list]], col: int) -> str:
    """Giá trị đầu tiên khác trống của cột (hiện làm ví dụ khi chọn cột)."""
    for _, v in rows[:50]:
        if col < len(v) and v[col] not in (None, ""):
            s = str(v[col]).strip()
            return s if len(s) <= 28 else s[:27] + "…"
    return ""


# ------------------------------------------------------------------ chuẩn hoá dữ liệu

def normalize_cccd(value) -> Tuple[str, str]:
    """(CCCD 12 số hoặc '', ghi chú). Sửa các lỗi hay gặp của Excel: mất số 0 đầu, dạng số mũ 5.6E+10."""
    if value is None or value == "":
        return "", "Trống số CCCD"
    if isinstance(value, float):
        value = int(value) if value.is_integer() else value
    s = str(value).strip()
    if re.fullmatch(r"\d+(\.\d+)?[eE]\+?\d+", s):
        try:
            s = str(int(float(s)))
        except ValueError:
            pass
    s = re.sub(r"\.0+$", "", s)
    digits = re.sub(r"\D", "", s)
    if len(digits) == 12:
        return digits, ""
    if len(digits) in (10, 11):
        padded = digits.zfill(12)
        if padded[:3] in R.PROVINCE_CODES:
            return padded, "Đã thêm số 0 đầu (Excel làm mất)"
    if len(digits) == 9:
        return "", f"Số CMND 9 số ({digits}), chưa phải CCCD"
    return "", f"Số CCCD không hợp lệ: “{s}”"


def name_key(name: str) -> str:
    return re.sub(r"[^a-z]", "", remove_accents(name or "").lower())


def pascal(name: str) -> str:
    """'Trần Văn Thái' / 'TRẦN VĂN THÁI' / 'TranVanThai' → 'TranVanThai'."""
    words = []
    for tok in re.findall(r"[A-Za-z]+", remove_accents(name or "")):
        if tok.isupper() or tok.islower():
            words.append(tok.capitalize())
        else:                                        # đã viết liền kiểu PascalCase → tách theo chữ hoa
            words += [w.capitalize() for w in re.findall(r"[A-Z][a-z]*|[a-z]+", tok)]
    return "".join(words)


def chibo_key(name: str) -> str:
    t = _norm(name)
    t = re.sub(r"\bchi bo\b", " ", t)
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def chibo_similarity(a: str, b: str) -> float:
    ka, kb = chibo_key(a), chibo_key(b)
    if not ka or not kb:
        return 0.0
    if ka == kb:
        return 1.0
    if ka in kb or kb in ka:
        return 0.9
    return SequenceMatcher(None, ka, kb).ratio()


def cccd_near(a: str, b: str) -> bool:
    """Lệch 1 chữ số hoặc 2 chữ số liền nhau bị đảo chỗ."""
    if len(a) != len(b) or a == b:
        return False
    diff = [i for i in range(len(a)) if a[i] != b[i]]
    if len(diff) == 1:
        return True
    return len(diff) == 2 and diff[1] == diff[0] + 1 and a[diff[0]] == b[diff[1]] and a[diff[1]] == b[diff[0]]


# ------------------------------------------------------------------ đối chiếu

@dataclass
class Person:
    row: int                           # dòng trong Excel
    cccd: str
    cccd_raw: str
    name: str
    chibo: str
    note: str = ""


@dataclass
class MatchRow:
    kind: str
    person: Optional[Person] = None
    dossier: Optional[R.Dossier] = None
    target_name: str = ""              # tên thư mục chuẩn theo Excel (CCCD_HoVaTen)
    target_chibo: str = ""             # thư mục chi bộ đúng theo Excel ('' = không xác định)
    notes: List[str] = field(default_factory=list)

    @property
    def need_rename(self) -> bool:
        return bool(self.dossier and self.person and self.person.cccd and self.target_name
                    and self.dossier.folder != self.target_name)

    @property
    def need_move(self) -> bool:
        return bool(self.dossier and self.target_chibo and self.dossier.chibo != self.target_chibo)

    @property
    def suggestion(self) -> str:
        out = []
        if self.kind == K_NO_FOLDER:
            out.append(f"Tạo thư mục {self.target_name}" + (f" trong “{self.target_chibo}”" if self.target_chibo else "")
                       + " rồi scan hồ sơ")
        elif self.kind == K_NOT_IN_LIST:
            out.append("Kiểm tra: sai số CCCD, đã chuyển đi hay thư mục thừa")
        else:
            if self.need_rename:
                out.append(f"Đổi tên thư mục → {self.target_name}")
            if self.need_move:
                out.append(f"Chuyển sang chi bộ “{self.target_chibo}”")
        return "; ".join(out)


@dataclass
class Comparison:
    rows: List[MatchRow]
    people: List[Person]
    skipped: List[Tuple[int, str]]     # (dòng, lý do) – dòng Excel không dùng được
    chibo_map: Dict[str, str]          # chi bộ trong Excel → thư mục chi bộ ghép được ('' = không có)
    report: R.Report
    table: Table
    columns: Dict[str, Optional[int]]

    def count(self, kind: str) -> int:
        return sum(1 for r in self.rows if r.kind == kind)


def extract_people(table: Table, cols: Dict[str, Optional[int]]) -> Tuple[List[Person], List[Tuple[int, str]]]:
    if cols.get("cccd") is None:
        raise CompareError("Chưa chọn cột Số CCCD.")
    if cols.get("name") is None:
        raise CompareError("Chưa chọn cột Họ và tên.")
    people, skipped = [], []
    for n, vals in table.rows:
        raw = vals[cols["cccd"]]
        name = str(vals[cols["name"]] or "").strip()
        chibo = str(vals[cols["chibo"]] or "").strip() if cols.get("chibo") is not None else ""
        if raw in (None, "") and not name:
            continue
        cccd, note = normalize_cccd(raw)
        if not cccd and not name:
            skipped.append((n, note))
            continue
        if not name:
            note = "; ".join(x for x in (note, "Trống họ tên") if x)
        people.append(Person(n, cccd, str(raw if raw is not None else ""), re.sub(r"\s+", " ", name), chibo, note))
    if not people:
        raise CompareError("Không đọc được người nào – kiểm tra lại dòng tiêu đề và các cột đã chọn.")
    return people, skipped


def _folder_cccd(d: R.Dossier) -> str:
    if d.cccd:
        return d.cccd
    m = re.match(r"\s*(\d{9,13})", d.folder)
    return m.group(1).zfill(12) if m and len(m.group(1)) in (10, 11, 12) else ""


def compare(table: Table, cols: Dict[str, Optional[int]], rep: R.Report) -> Comparison:
    people, skipped = extract_people(table, cols)
    folders_chibo = list(rep.chibos)

    # ghép chi bộ Excel ↔ thư mục chi bộ (gần đúng)
    chibo_map: Dict[str, str] = {}
    for c in sorted({p.chibo for p in people if p.chibo}):
        best, score = "", 0.0
        for f in folders_chibo:
            s = chibo_similarity(c, f)
            if s > score:
                best, score = f, s
        chibo_map[c] = best if score >= 0.6 else ""

    by_cccd: Dict[str, List[R.Dossier]] = {}
    for d in rep.dossiers:
        k = _folder_cccd(d)
        if k:
            by_cccd.setdefault(k, []).append(d)
    seen_excel: Dict[str, List[Person]] = {}
    for p in people:
        if p.cccd:
            seen_excel.setdefault(p.cccd, []).append(p)

    rows: List[MatchRow] = []
    used: set = set()
    unmatched_people: List[MatchRow] = []
    for p in people:
        target = f"{p.cccd}_{pascal(p.name)}" if p.cccd and pascal(p.name) else ""
        tchibo = chibo_map.get(p.chibo, "") if p.chibo else ""
        row = MatchRow(K_NO_FOLDER, person=p, target_name=target, target_chibo=tchibo)
        if p.note:
            row.notes.append(p.note)
        if p.cccd and len(seen_excel.get(p.cccd, [])) > 1:
            others = ", ".join(str(x.row) for x in seen_excel[p.cccd] if x is not p)
            row.notes.append(f"Trùng CCCD trong Excel (dòng {others})")
        if p.chibo and not tchibo:
            row.notes.append(f"Chi bộ “{p.chibo}” chưa có thư mục tương ứng")
        cands = [d for d in by_cccd.get(p.cccd, []) if id(d) not in used] if p.cccd else []
        if cands:
            d = next((x for x in cands if name_key(x.name) == name_key(p.name)), cands[0])
            used.add(id(d))
            row.dossier = d
            _judge(row)
            if len(by_cccd[p.cccd]) > 1:
                row.notes.append("Có nhiều thư mục cùng số CCCD: " + ", ".join(f"{x.chibo}/{x.folder}"
                                                                           for x in by_cccd[p.cccd]))
            rows.append(row)
        else:
            unmatched_people.append(row)

    leftovers = [d for d in rep.dossiers if id(d) not in used]
    # gần đúng: CCCD lệch 1 số / đảo 2 số
    for row in unmatched_people:
        p = row.person
        if p.cccd:
            near = [d for d in leftovers if id(d) not in used and cccd_near(p.cccd, _folder_cccd(d))]
            same = [d for d in near if name_key(d.name) == name_key(p.name)]
            if same:
                d = same[0]
                used.add(id(d))
                row.dossier, row.kind = d, K_TYPO
                row.notes.append(f"CCCD thư mục {_folder_cccd(d)} khác Excel {p.cccd} (lệch 1–2 số), họ tên giống")
                if row.target_chibo and d.chibo != row.target_chibo:
                    row.notes.append(f"Đang ở chi bộ “{d.chibo}”")
            else:
                for d in near:
                    row.notes.append(f"CCCD gần giống thư mục {d.chibo}/{d.folder} nhưng họ tên khác")
        rows.append(row)

    for d in leftovers:
        if id(d) in used:
            continue
        row = MatchRow(K_NOT_IN_LIST, dossier=d)
        k = _folder_cccd(d)
        if not k:
            row.notes.append("Tên thư mục không có số CCCD hợp lệ")
        for p in people:
            if k and p.cccd and cccd_near(k, p.cccd):
                row.notes.append(f"CCCD gần giống Excel dòng {p.row} ({p.cccd} – {p.name})")
        rows.append(row)

    order = {c: i for i, c in enumerate(folders_chibo)}
    rows.sort(key=lambda r: (KIND_ORDER.index(r.kind) if r.kind != K_MATCH else 99,
                             order.get(r.dossier.chibo if r.dossier else r.target_chibo, 9999),
                             r.person.row if r.person else 10 ** 9))
    return Comparison(rows, people, skipped, chibo_map, rep, table, dict(cols))


def _judge(row: MatchRow):
    p, d = row.person, row.dossier
    name_ok = name_key(d.name) == name_key(p.name) if p.name else True
    if not p.chibo:
        chibo_ok = True
    elif row.target_chibo:                       # đã ghép được thư mục chi bộ → so chính xác
        chibo_ok = row.target_chibo == d.chibo
    else:
        chibo_ok = chibo_similarity(p.chibo, d.chibo) >= 0.6
    if not name_ok:
        row.kind = K_NAME
        row.notes.append(f"Tên thư mục “{d.name or d.folder}” ≠ Excel “{p.name}”")
        if not chibo_ok:
            row.notes.append(f"Khác chi bộ: thư mục ở “{d.chibo}”, Excel ghi “{p.chibo}”")
    elif not chibo_ok:
        row.kind = K_CHIBO
        row.notes.append(f"Thư mục ở “{d.chibo}”, Excel ghi “{p.chibo}”")
    else:
        row.kind = K_MATCH
        if row.target_name and d.folder != row.target_name:
            row.notes.append("Tên thư mục chưa đúng chuẩn")


def summary_by_chibo(cmp: Comparison) -> List[dict]:
    """Mỗi chi bộ (theo Excel; chi bộ chỉ có thư mục đứng cuối)."""
    groups: Dict[str, dict] = {}

    def g(key: str, folder: str) -> dict:
        if key not in groups:
            groups[key] = {"chibo": key, "folder": folder, "excel": 0, "has_folder": 0, "match": 0, "fix": 0,
                           "no_folder": 0, "not_in_list": 0, "full": 0, "lack": 0}
        return groups[key]

    for r in cmp.rows:
        if r.person:
            x = g(r.person.chibo or "(không ghi chi bộ)", r.target_chibo)
            x["excel"] += 1
        else:
            x = g(f"(chỉ có thư mục) {r.dossier.chibo}", r.dossier.chibo)
        if r.dossier:
            if r.person:
                x["has_folder"] += 1
            x["full" if r.dossier.complete else "lack"] += 1
        x["match"] += r.kind == K_MATCH
        x["fix"] += r.kind in NEED_FIX
        x["no_folder"] += r.kind == K_NO_FOLDER
        x["not_in_list"] += r.kind == K_NOT_IN_LIST
    return sorted(groups.values(), key=lambda x: (x["chibo"].startswith("("), R.natural_key(x["chibo"])))


# ------------------------------------------------------------------ thao tác sửa (có hoàn tác)

@dataclass
class Op:
    kind: str                          # rename | move | mkdir
    src: str
    dst: str
    ok: bool = False
    message: str = ""


def _safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', " ", name).strip().rstrip(".")


def plan_renames(cmp: Comparison) -> List[Tuple[str, str]]:
    return [(r.dossier.path, r.target_name) for r in cmp.rows if r.kind != K_NOT_IN_LIST and r.need_rename]


def plan_moves(cmp: Comparison) -> List[Tuple[str, str]]:
    root = cmp.report.root if cmp.report.mode == "root" else ""
    if not root:
        return []
    return [(r.dossier.path, os.path.join(root, r.target_chibo, os.path.basename(r.dossier.path)))
            for r in cmp.rows if r.kind != K_NOT_IN_LIST and r.need_move]


def plan_creates(cmp: Comparison) -> List[str]:
    rep = cmp.report
    out = []
    for r in cmp.rows:
        if r.kind != K_NO_FOLDER or not r.target_name:
            continue
        if rep.mode == "root":
            chibo = r.target_chibo or (_safe_name(r.person.chibo) if r.person.chibo else "")
            if not chibo:
                continue
            out.append(os.path.join(rep.root, chibo, r.target_name))
        elif rep.mode == "chibo":
            out.append(os.path.join(rep.root, r.target_name))
    return out


def apply_ops(renames=(), moves=(), creates=()) -> List[Op]:
    ops: List[Op] = []
    for src, new_name in renames:
        ops.append(_move(Op("rename", src, os.path.join(os.path.dirname(src), new_name))))
    for src, dst in moves:
        ops.append(_move(Op("move", src, dst)))
    for path in creates:
        op = Op("mkdir", "", path)
        if os.path.exists(path):
            op.message = "Đã có thư mục"
        else:
            try:
                parent = os.path.dirname(path)
                made_parent = not os.path.isdir(parent)
                os.makedirs(path)
                op.ok, op.message = True, "Đã tạo"
                op.src = parent if made_parent else ""      # nhớ thư mục chi bộ mới tạo để hoàn tác
            except OSError as e:
                op.message = f"Không tạo được: {e.strerror or e}"
        ops.append(op)
    return ops


def _move(op: Op) -> Op:
    src, dst = op.src, op.dst
    if not os.path.isdir(src):
        op.message = "Thư mục không còn tồn tại"
        return op
    same = os.path.normcase(os.path.abspath(src)) == os.path.normcase(os.path.abspath(dst))
    if os.path.exists(dst) and not same:
        op.message = "Đã có thư mục trùng tên ở nơi đến – bỏ qua"
        return op
    try:
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        if same:
            tmp = src + ".~doiten"
            os.rename(src, tmp)
            os.rename(tmp, dst)
        else:
            os.rename(src, dst)
        op.ok, op.message = True, "Xong"
    except OSError as e:
        op.message = f"Không thực hiện được: {e.strerror or e}"
    return op


def undo_ops(ops: Sequence[Op]) -> List[Op]:
    """Hoàn tác theo thứ tự ngược; thư mục mới tạo chỉ xoá khi vẫn còn trống."""
    out: List[Op] = []
    for op in reversed([o for o in ops if o.ok]):
        if op.kind == "mkdir":
            u = Op("rmdir", op.dst, "")
            try:
                os.rmdir(op.dst)
                if op.src:
                    os.rmdir(op.src)
                u.ok, u.message = True, "Đã xoá thư mục trống"
            except OSError:
                u.message = "Thư mục đã có file – giữ lại"
            out.append(u)
        else:
            out.append(_move(Op("undo", op.dst, op.src)))
    return out
