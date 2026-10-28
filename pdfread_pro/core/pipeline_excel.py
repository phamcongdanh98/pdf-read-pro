"""Xuất Excel kết quả quy trình 1 click: tổng hợp từng bước + chi tiết từng file."""
from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook

from . import pipeline as P
from .report_excel import CENTER, FILL_NO, FILL_NOTE, FILL_OK, FILL_TOTAL, WRAP, _font, _header, _row, _setup, _title

LEVEL_FILL = {P.OK: FILL_OK, P.WARN: FILL_NOTE, P.FAIL: FILL_NO, P.SKIP: None, P.CANCEL: FILL_NOTE}
LEVEL_TEXT = {P.OK: "Đạt", P.WARN: "Cần xem", P.FAIL: "Lỗi", P.SKIP: "Không chạy", P.CANCEL: "Đã dừng"}


def export(path: str, rep: P.Report) -> str:
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    if not os.path.isdir(os.path.dirname(os.path.abspath(path))):
        raise OSError("Thư mục lưu không tồn tại.")
    cfg = rep.cfg
    sub = (f"Hồ sơ gốc: {cfg.source}   ·   Quy tắc đổi tên: {cfg.rule_label or '—'}   ·   Giới hạn: {cfg.limit_kb} KB/trang   ·   "
           f"Ngày lập: {datetime.now():%d/%m/%Y %H:%M}")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    cols = ["Bước", "Tên bước", "Kết quả", "Chi tiết", "Thời gian (giây)"]
    _title(ws, "Kết quả quy trình 1 click", sub, len(cols))
    _header(ws, 3, cols, [8, 24, 14, 90, 16])
    r = 4
    for i, s in enumerate(rep.stages):
        r = _row(ws, r, [i, s.name, LEVEL_TEXT.get(s.status, s.status), s.summary, round(s.seconds, 1)],
                 [CENTER, WRAP, CENTER, WRAP, CENTER])
        f = LEVEL_FILL.get(s.status)
        if f:
            ws.cell(row=r - 1, column=3).fill = f
    r = _row(ws, r, ["", "Thư mục làm việc", "", cfg.work, ""], [CENTER, WRAP, CENTER, WRAP, CENTER])
    r = _row(ws, r, ["", "Thư mục gom (để ký)", "", cfg.gather, ""], [CENTER, WRAP, CENTER, WRAP, CENTER])
    if rep.error:
        r = _row(ws, r, ["", "Ghi chú", "", rep.error, ""], [CENTER, WRAP, CENTER, WRAP, CENTER])
    r = _row(ws, r, ["", "TỔNG", f"{len(rep.recs)} file", f"Giảm {rep.saved_bytes / 2 ** 20:,.1f} MB · đã gom {rep.gathered} file",
                     round(rep.seconds, 1)], [CENTER, WRAP, CENTER, WRAP, CENTER])
    for c in range(1, len(cols) + 1):
        ws.cell(row=r - 1, column=c).fill = FILL_TOTAL
        ws.cell(row=r - 1, column=c).font = _font(True)
    _setup(ws, 3, r - 2, len(cols), landscape=True, filter_=False)

    ws = wb.create_sheet("Chi tiết")
    cols = ["STT", "Hồ sơ", "Tên gốc", "Tên mới", "Đổi tên", "Dung lượng trước (KB)", "Dung lượng sau (KB)", "Nén",
            "Gom để ký", "Kết luận"]
    _title(ws, "Kết quả quy trình 1 click – chi tiết từng file", sub, len(cols))
    _header(ws, 3, cols, [7, 30, 40, 40, 38, 14, 14, 44, 36, 12])
    r = 4
    for i, rec in enumerate(rep.recs, 1):
        r = _row(ws, r, [i, rec.folder, rec.old_name, rec.new_name or rec.old_name, rec.rename, round(rec.before / 1024),
                         round(rec.after / 1024), rec.comp, rec.gather, LEVEL_TEXT[rec.level]],
                 [CENTER, WRAP, WRAP, WRAP, WRAP, CENTER, CENTER, WRAP, WRAP, CENTER])
        f = LEVEL_FILL.get(rec.level)
        if f:
            ws.cell(row=r - 1, column=10).fill = f
    _setup(ws, 3, r - 1, len(cols))
    wb.properties.title = "Kết quả quy trình 1 click"
    wb.properties.creator = "PDF-Read Pro"
    tmp = path + ".partial"
    wb.save(tmp)
    os.replace(tmp, path)
    return path
