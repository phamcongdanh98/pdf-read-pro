"""Xuất báo cáo hồ sơ số hoá ra Excel (.xlsx).

Bản chính (dễ đọc, dễ in):
    1. Tổng hợp chi bộ   – mỗi chi bộ: tổng hồ sơ, Đủ, Thiếu, thiếu từng mục 01…05, tỉ lệ đủ
    2. Danh sách toàn bộ – mỗi người 1 dòng: chi bộ, CCCD, họ tên, 01…05 Có/Thiếu, Kết quả, Còn thiếu, Ghi chú
    3. Danh sách thiếu   – chỉ những người còn thiếu, có cột "Đã bổ sung" để đánh dấu
    4. Giải thích        – cấu trúc thư mục chuẩn và ý nghĩa các cột
Kèm báo cáo nâng cao (tuỳ chọn): Chi tiết file · Thống kê giấy tờ (cả bộ quy tắc) · Vấn đề cần xử lý."""
from __future__ import annotations

import os
import re
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import report as R

NAVY = "1E3A8A"
FONT = "Times New Roman"
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_DOC = PatternFill("solid", fgColor="1D4ED8")
FILL_MISS_HEAD = PatternFill("solid", fgColor="991B1B")
FILL_TOTAL = PatternFill("solid", fgColor="E2E8F0")
FILL_GROUP = PatternFill("solid", fgColor="EFF6FF")
FILL_OK = PatternFill("solid", fgColor="DCFCE7")
FILL_NO = PatternFill("solid", fgColor="FEE2E2")
FILL_NOTE = PatternFill("solid", fgColor="FEF9C3")
THIN = Side(style="thin", color="94A3B8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
GREEN, RED, AMBER = "15803D", "B91C1C", "92400E"


# ------------------------------------------------------------------ tiện ích định dạng

def _font(bold=False, color="000000", size=12, **kw) -> Font:
    return Font(name=FONT, bold=bold, color=color, size=size, **kw)


def _title(ws: Worksheet, text: str, sub: str, width: int):
    for r, (t, f) in enumerate(((text, _font(True, NAVY, 15)), (sub, _font(False, "475569", 11, italic=True))), 1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        c = ws.cell(row=r, column=1, value=t)
        c.font = f
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24


def _header(ws: Worksheet, row: int, cols: List[str], widths: List[float], fills: Dict[int, PatternFill] = None):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = _font(True, "FFFFFF", 11)
        cell.fill = (fills or {}).get(i, FILL_HEAD)
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 48


def _row(ws: Worksheet, row: int, values: list, aligns: List[Alignment]) -> int:
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = _font()
        cell.border = BORDER
        cell.alignment = aligns[i - 1]
    return row + 1


def _setup(ws: Worksheet, head_row: int, last_row: int, ncols: int, landscape=True, filter_=True):
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    if filter_ and last_row > head_row:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(ncols)}{last_row}"
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.oddFooter.center.text = "Trang &P / &N"


def _link(cell, path: str):
    cell.hyperlink = "file:///" + path.replace("\\", "/")
    cell.font = _font(color="1D4ED8", underline="single")


def _code_key(c: str):
    m = re.match(r"(\d+)(.*)", c)
    return (int(m.group(1)), m.group(2)) if m else (9999, c)


def _given_name_key(d: R.Dossier):
    """Sắp theo tên (từ cuối) rồi họ đệm – kiểu danh sách Việt Nam."""
    words = d.name.split()
    return (words[-1].lower() if words else "~", d.name.lower(), d.cccd)


def _missing_text(d: R.Dossier, rep: R.Report) -> str:
    return "; ".join(f"{c} – {R.doc_title(c, rep.rules)}{R.missing_hint(d, c, rep.sign_filter)}" for c in d.missing)


def _people(rep: R.Report) -> List[R.Dossier]:
    order = {c: i for i, c in enumerate(rep.chibos)}
    return sorted(rep.dossiers, key=lambda d: (order.get(d.chibo, 9999), _given_name_key(d)))


# ------------------------------------------------------------------ 1. Tổng hợp chi bộ

def _sheet_summary(wb: Workbook, rep: R.Report, sub: str):
    ws = wb.active
    ws.title = "Tổng hợp chi bộ"
    req = list(rep.required)
    cols = (["STT", "Chi bộ", "Tổng số hồ sơ", "Đủ", "Thiếu"] + [f"Thiếu {c}\n{R.doc_title(c, rep.rules)}" for c in req]
            + ["Tỉ lệ đủ"])
    n = len(cols)
    _title(ws, "BÁO CÁO TỔNG HỢP HỒ SƠ SỐ HÓA THEO CHI BỘ", sub, n)
    total = len(rep.dossiers)
    full = sum(1 for d in rep.dossiers if d.complete)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n)
    c = ws.cell(row=3, column=1, value=f"Tổng cộng: {len(rep.chibos)} chi bộ · {total} hồ sơ · Đủ: {full} · "
                                       f"Thiếu: {total - full}")
    c.font = _font(True, NAVY, 12)
    c.alignment = Alignment(horizontal="center")
    head = 5
    fills = {4: PatternFill("solid", fgColor=GREEN), 5: PatternFill("solid", fgColor=RED)}
    fills.update({6 + i: FILL_MISS_HEAD for i in range(len(req))})
    _header(ws, head, cols, [6, 40, 10, 9, 9] + [14] * len(req) + [10], fills)
    r = head + 1
    for i, cb in enumerate(rep.chibos, 1):
        ds = [d for d in rep.dossiers if d.chibo == cb]
        ok = sum(1 for d in ds if d.complete)
        miss = [sum(1 for d in ds if c in d.missing) for c in req]
        r = _row(ws, r, [i, cb, len(ds), ok, len(ds) - ok] + miss + [ok / len(ds) if ds else 0],
                 [CENTER, WRAP] + [CENTER] * (n - 2))
        ws.cell(row=r - 1, column=n).number_format = "0%"
        ws.cell(row=r - 1, column=4).font = _font(True, GREEN)
        if len(ds) - ok:
            ws.cell(row=r - 1, column=5).font = _font(True, RED)
        for j, v in enumerate(miss):
            if v:
                cell = ws.cell(row=r - 1, column=6 + j)
                cell.fill = FILL_NO
                cell.font = _font(True, RED)
    last = r - 1
    if last > head:
        L = get_column_letter
        tot = ["", "TỔNG CỘNG"] + [f"=SUM({L(c)}{head + 1}:{L(c)}{last})" for c in range(3, n)]
        tot += [f"=IF(C{r}=0,0,D{r}/C{r})"]
        _row(ws, r, tot, [CENTER, WRAP] + [CENTER] * (n - 2))
        for cidx in range(1, n + 1):
            ws.cell(row=r, column=cidx).fill = FILL_TOTAL
            ws.cell(row=r, column=cidx).font = _font(True)
        ws.cell(row=r, column=n).number_format = "0%"
    _setup(ws, head, last, n)
    r += 2
    extra = list(rep.notes)
    if rep.stray:
        extra.append(f"Có {len(rep.stray)} file PDF nằm ngoài thư mục hồ sơ: "
                     + "; ".join(s.rel for s in rep.stray[:10]) + (" …" if len(rep.stray) > 10 else ""))
    for t in extra:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=n)
        ws.cell(row=r, column=2, value="• " + t).font = _font(False, AMBER, 11)
        r += 1


# ------------------------------------------------------------------ 2 + 3. Danh sách người

def _sheet_people(wb: Workbook, rep: R.Report, sub: str, only_missing: bool):
    req = list(rep.required)
    title = "DANH SÁCH HỒ SƠ CÒN THIẾU" if only_missing else "DANH SÁCH TOÀN BỘ HỒ SƠ"
    ws = wb.create_sheet("Danh sách thiếu" if only_missing else "Danh sách toàn bộ")
    people = [d for d in _people(rep) if not only_missing or not d.complete]
    cols = (["STT", "Chi bộ", "Số CCCD", "Họ và tên"] + [f"{c}\n{R.doc_title(c, rep.rules)}" for c in req]
            + ["Giấy tờ khác", "Tổng số file", "Tổng số trang", "Kết quả", "Còn thiếu", "Ghi chú"])
    if only_missing:
        cols.append("Đã bổ sung")
    n = len(cols)
    full = sum(1 for d in rep.dossiers if d.complete)
    _title(ws, title, sub + (f"  ·  {len(people)} hồ sơ còn thiếu" if only_missing else
                             f"  ·  {len(people)} hồ sơ: Đủ {full} · Thiếu {len(people) - full}"), n)
    head = 4
    fills = {5 + i: FILL_DOC for i in range(len(req))}
    widths = [6, 34, 15, 26] + [13] * len(req) + [9, 8, 8, 10, 36, 40] + ([10] if only_missing else [])
    _header(ws, head, cols, widths, fills)
    c_res = 5 + len(req) + 3                     # cột Kết quả
    r = head + 1
    for i, d in enumerate(people, 1):
        marks = ["Có" if d.have.get(c) else "Thiếu" for c in req]
        vals = ([i, d.chibo, d.cccd or "(chưa rõ)", d.name or d.folder] + marks
                + [d.other_docs, len(d.files), d.pages, d.result, _missing_text(d, rep), "\n".join(d.notes)])
        if only_missing:
            vals.append("☐")
        r = _row(ws, r, vals, [CENTER, WRAP, CENTER, WRAP] + [CENTER] * len(req)
                 + [CENTER, CENTER, CENTER, CENTER, WRAP, WRAP] + ([CENTER] if only_missing else []))
        ws.cell(row=r - 1, column=3).number_format = "@"
        _link(ws.cell(row=r - 1, column=4), d.path)
        for j, m in enumerate(marks):
            cell = ws.cell(row=r - 1, column=5 + j)
            if m == "Thiếu":
                cell.fill = FILL_NO
                cell.font = _font(True, RED)
            else:
                cell.font = _font(False, GREEN)
        res = ws.cell(row=r - 1, column=c_res)
        res.fill = FILL_OK if d.complete else FILL_NO
        res.font = _font(True, GREEN if d.complete else RED)
        if not d.complete:
            ws.cell(row=r - 1, column=c_res + 1).font = _font(True, RED)
        if d.notes:
            ws.cell(row=r - 1, column=c_res + 2).fill = FILL_NOTE
    if not people:
        ws.cell(row=r, column=2, value="Không có hồ sơ nào còn thiếu." if only_missing else "Không có hồ sơ.").font = \
            _font(True, GREEN)
    _setup(ws, head, r - 1, n)
    return ws


# ------------------------------------------------------------------ 4. Giải thích

def _sheet_guide(wb: Workbook, rep: R.Report):
    ws = wb.create_sheet("Giải thích")
    _title(ws, "GIẢI THÍCH BÁO CÁO", "Cách đọc các cột và quy chuẩn sắp xếp hồ sơ", 2)
    req = list(rep.required)
    rows = [
        ("Cấu trúc thư mục", "Hồ sơ số hóa  →  <Chi bộ>  →  <Số CCCD>_<HoVaTen>  →  các file PDF"),
        ("Thư mục hồ sơ", "12 số CCCD + dấu “_” + họ tên không dấu, viết hoa chữ cái đầu mỗi từ, không khoảng trắng. "
                          "VD: 056123456789_TranVanThai"),
        ("Tên file", "Đặt theo bộ quy tắc 1361: <mã>.<Ten_mau>.pdf, VD 01.Ly_lich_nguoi_xin_vao_Dang.pdf. "
                     "Giấy tờ nhiều phần thêm .1, .2; file đã ký số có thể thêm .signed"),
        ("File bắt buộc", "; ".join(f"{c} – {R.doc_title(c, rep.rules)}" for c in req)),
        ("Cột 01…05", "“Có” = hồ sơ có file đó · “Thiếu” = chưa có file đó (ô màu đỏ). "
                      "File có nhưng đặt sai tên vẫn tính là “Có”, kèm nhắc ở cột Ghi chú"),
        ("Kết quả", "“Đủ” = có đủ tất cả file bắt buộc · “Thiếu” = thiếu ít nhất 1 file bắt buộc"),
        ("Còn thiếu", "Liệt kê các file bắt buộc còn thiếu của người đó"),
        ("Giấy tờ khác", "Số file thuộc bộ quy tắc 1361 nhưng không bắt buộc (mã 06 trở đi) – không ảnh hưởng kết quả"),
        ("Ghi chú", "Vấn đề khác cần xem lại (không làm hồ sơ thành “Thiếu”): thư mục hoặc file đặt sai tên, "
                    "trùng số CCCD, file hỏng, file chưa rõ loại, thư mục con…"),
        ("Họ và tên", "Lấy từ tên thư mục (không dấu)"),
        ("Tính file", f"Báo cáo này tính: {R.SIGN_LABELS[rep.sign_filter]}. Khi ký số thường có 2 bản: X.pdf (chưa ký) "
                      "và X.signed.pdf (đã ký). “Chỉ file đã ký”: mục chỉ có bản chưa ký vẫn tính Thiếu, ghi "
                      "“(mới có bản chưa ký)”"),
        ("Mẹo", "Bấm vào họ tên để mở thư mục hồ sơ. Dùng nút lọc ở dòng tiêu đề để lọc theo Chi bộ / Kết quả / "
                "từng cột 01…05"),
    ]
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    r = 4
    for k, v in rows:
        a = ws.cell(row=r, column=1, value=k)
        a.font = _font(True, NAVY)
        a.alignment = WRAP
        a.border = BORDER
        b = ws.cell(row=r, column=2, value=v)
        b.font = _font()
        b.alignment = WRAP
        b.border = BORDER
        ws.row_dimensions[r].height = 36
        r += 1
    _setup(ws, 3, r - 1, 2, landscape=False, filter_=False)
    ws.freeze_panes = None


# ------------------------------------------------------------------ nâng cao

def _sheet_files(wb: Workbook, rep: R.Report, sub: str):
    ws = wb.create_sheet("Chi tiết file")
    cols = ["STT", "Chi bộ", "Số CCCD", "Họ và tên", "Tên file", "Mã", "Loại giấy tờ", "Bắt buộc", "Đúng tên mẫu",
            "Số trang", "Dung lượng (KB)", "Chữ ký số", "Ghi chú"]
    _title(ws, "CHI TIẾT TỪNG FILE", sub, len(cols))
    head = 4
    _header(ws, head, cols, [6, 30, 15, 22, 50, 7, 34, 9, 9, 8, 11, 12, 34])
    sig = {"ok": "Còn nguyên", "modified": "Sửa sau ký", "broken": "Hỏng", "none": "Chưa ký", "unreadable": "Không đọc được"}
    r, k = head + 1, 0
    for d in _people(rep):
        for f in d.files:
            k += 1
            r = _row(ws, r, [k, d.chibo, d.cccd, d.name, f.fi.name, f.code, R.doc_title(f.code, rep.rules) if f.code else "",
                             "Có" if f.required else "", ("Đúng" if f.name_ok else "Sai") if f.code else "",
                             f.fi.pages if f.fi.pages is not None else "", round(f.fi.size / 1024, 1),
                             sig.get(f.fi.sig_status, f.fi.sig_status), f.note],
                     [CENTER, WRAP, CENTER, WRAP, WRAP, CENTER, WRAP, CENTER, CENTER, CENTER, CENTER, CENTER, WRAP])
            ws.cell(row=r - 1, column=3).number_format = "@"
            ws.cell(row=r - 1, column=6).number_format = "@"
            _link(ws.cell(row=r - 1, column=5), f.fi.path)
            if f.code and not f.name_ok:
                ws.cell(row=r - 1, column=9).fill = FILL_NOTE
    _setup(ws, head, r - 1, len(cols))


def _sheet_doc_stats(wb: Workbook, rep: R.Report, sub: str):
    ws = wb.create_sheet("Thống kê giấy tờ")
    cols = ["Mã", "Tên file mẫu", "Loại giấy tờ", "Bắt buộc", "Số hồ sơ có", "Tỉ lệ", "Số file", "Số trang"]
    _title(ws, "THỐNG KÊ THEO LOẠI GIẤY TỜ (BỘ QUY TẮC 1361)", sub, len(cols))
    head = 4
    _header(ws, head, cols, [8, 60, 46, 9, 11, 9, 9, 10])
    per: Dict[str, list] = {c: [set(), 0, 0] for c in rep.rules}
    for d in rep.dossiers:
        for f in d.files:
            if f.code in per:
                per[f.code][0].add(d.path)
                per[f.code][1] += 1
                per[f.code][2] += f.fi.pages or 0
    total = len(rep.dossiers)
    r = head + 1
    for c in sorted(rep.rules, key=_code_key):
        has, nf, npg = per[c]
        r = _row(ws, r, [c, rep.rules[c] + ".pdf", R.doc_title(c, rep.rules), "Có" if c in rep.required else "",
                         len(has), len(has) / total if total else 0, nf, npg],
                 [CENTER, WRAP, WRAP, CENTER, CENTER, CENTER, CENTER, CENTER])
        ws.cell(row=r - 1, column=1).number_format = "@"
        ws.cell(row=r - 1, column=6).number_format = "0.0%"
        if c in rep.required:
            for k in range(1, len(cols) + 1):
                ws.cell(row=r - 1, column=k).fill = FILL_GROUP
    _setup(ws, head, r - 1, len(cols))


def _sheet_issues(wb: Workbook, rep: R.Report, sub: str):
    ws = wb.create_sheet("Vấn đề cần xử lý")
    cols = ["STT", "Mức độ", "Chi bộ", "Hồ sơ / File", "Vấn đề", "Cách xử lý", "Đã xử lý"]
    issues = rep.all_issues()
    _title(ws, "DANH SÁCH VẤN ĐỀ CẦN XỬ LÝ", sub, len(cols))
    head = 4
    _header(ws, head, cols, [6, 10, 30, 32, 60, 50, 10])
    r = head + 1
    for i, (cb, folder, iss) in enumerate(issues, 1):
        r = _row(ws, r, [i, iss.level, cb, folder, iss.text, iss.fix, "☐"],
                 [CENTER, CENTER, WRAP, WRAP, WRAP, WRAP, CENTER])
        ws.cell(row=r - 1, column=2).font = _font(True, RED if iss.level == R.LV_ERROR else AMBER)
    _setup(ws, head, r - 1, len(cols))


# ------------------------------------------------------------------ xuất

def export_excel(rep: R.Report, path: str, advanced: bool = False) -> str:
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    folder = os.path.dirname(os.path.abspath(path))
    if not os.path.isdir(folder):
        raise OSError("Thư mục lưu không tồn tại.")
    sub = (f"Thư mục: {rep.root}   ·   Tính file: {R.SIGN_LABELS[rep.sign_filter]}   ·   "
           f"Ngày lập: {rep.created:%d/%m/%Y %H:%M}")
    wb = Workbook()
    _sheet_summary(wb, rep, sub)
    _sheet_people(wb, rep, sub, only_missing=False)
    _sheet_people(wb, rep, sub, only_missing=True)
    _sheet_guide(wb, rep)
    if advanced:
        _sheet_files(wb, rep, sub)
        _sheet_doc_stats(wb, rep, sub)
        _sheet_issues(wb, rep, sub)
    wb.properties.title = "Báo cáo hồ sơ số hóa"
    wb.properties.creator = "PDF-Read Pro"
    tmp = path + ".partial"
    wb.save(tmp)
    os.replace(tmp, path)
    return path
