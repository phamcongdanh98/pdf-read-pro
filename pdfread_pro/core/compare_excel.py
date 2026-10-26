"""Xuất kết quả ĐỐI CHIẾU danh sách hệ thống ↔ thư mục hồ sơ ra Excel:
Tổng hợp chi bộ · Đối chiếu chi tiết · Chưa có thư mục · Ngoài danh sách · Cần sửa · Giải thích."""
from __future__ import annotations

import os
from typing import List

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from . import compare as C, report as R
from .report_excel import (CENTER, FILL_NO, FILL_NOTE, FILL_OK, FILL_TOTAL, GREEN, RED, AMBER, WRAP, _font,
                           _header, _link, _row, _setup, _title)

KIND_FILL = {C.K_MATCH: FILL_OK, C.K_NAME: FILL_NOTE, C.K_CHIBO: FILL_NOTE, C.K_TYPO: FILL_NOTE,
             C.K_NO_FOLDER: FILL_NO, C.K_NOT_IN_LIST: FILL_NO}
KIND_COLOR = {C.K_MATCH: GREEN, C.K_NAME: AMBER, C.K_CHIBO: AMBER, C.K_TYPO: AMBER,
              C.K_NO_FOLDER: RED, C.K_NOT_IN_LIST: RED}


def _detail_sheet(wb: Workbook, cmp: C.Comparison, title: str, sheet: str, rows: List[C.MatchRow], sub: str,
                  check_col: bool = False):
    ws = wb.create_sheet(sheet)
    req = list(cmp.report.required)
    cols = (["STT", "Dòng Excel", "Số CCCD (Excel)", "Họ và tên (Excel)", "Chi bộ (Excel)", "Thư mục hồ sơ",
             "Chi bộ (thư mục)", "Kết quả đối chiếu"] + req + ["Hồ sơ", "Còn thiếu", "Gợi ý xử lý", "Ghi chú"]
            + (["Đã xử lý"] if check_col else []))
    n = len(cols)
    _title(ws, title, sub + f"  ·  {len(rows)} dòng", n)
    head = 4
    _header(ws, head, cols, [6, 8, 15, 26, 28, 32, 28, 14] + [6] * len(req) + [8, 22, 40, 44]
            + ([9] if check_col else []))
    c_kind = 8
    r = head + 1
    for i, m in enumerate(rows, 1):
        p, d = m.person, m.dossier
        marks = [("Có" if d.have.get(c) else "Thiếu") if d else "" for c in req]
        vals = ([i, p.row if p else "", (p.cccd or p.cccd_raw) if p else "", p.name if p else "",
                 p.chibo if p else "", d.folder if d else "", d.chibo if d else "", m.kind] + marks
                + [d.result if d else "", ", ".join(d.missing) if d else "", m.suggestion, "\n".join(m.notes)]
                + (["☐"] if check_col else []))
        r = _row(ws, r, vals, [CENTER, CENTER, CENTER, WRAP, WRAP, WRAP, WRAP, CENTER] + [CENTER] * len(req)
                 + [CENTER, WRAP, WRAP, WRAP] + ([CENTER] if check_col else []))
        ws.cell(row=r - 1, column=3).number_format = "@"
        if d:
            _link(ws.cell(row=r - 1, column=6), d.path)
        k = ws.cell(row=r - 1, column=c_kind)
        k.fill = KIND_FILL[m.kind]
        k.font = _font(True, KIND_COLOR[m.kind])
        for j, mk in enumerate(marks):
            cell = ws.cell(row=r - 1, column=c_kind + 1 + j)
            if mk == "Thiếu":
                cell.fill = FILL_NO
                cell.font = _font(True, RED)
            elif mk:
                cell.font = _font(False, GREEN)
        if d:
            res = ws.cell(row=r - 1, column=c_kind + 1 + len(req))
            res.font = _font(True, GREEN if d.complete else RED)
        if m.notes:
            ws.cell(row=r - 1, column=n - (1 if check_col else 0)).fill = FILL_NOTE
    if not rows:
        ws.cell(row=r, column=2, value="Không có dòng nào.").font = _font(True, GREEN)
    _setup(ws, head, r - 1, n)


def export_compare(cmp: C.Comparison, path: str) -> str:
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    folder = os.path.dirname(os.path.abspath(path))
    if not os.path.isdir(folder):
        raise OSError("Thư mục lưu không tồn tại.")
    rep = cmp.report
    sub = (f"Danh sách: {os.path.basename(cmp.table.path)} ({cmp.table.sheet})   ·   Thư mục: {rep.root}   ·   "
           f"Tính file: {R.SIGN_LABELS[rep.sign_filter]}   ·   Ngày lập: {rep.created:%d/%m/%Y %H:%M}")
    wb = Workbook()

    # ---- 1. Tổng hợp chi bộ
    ws = wb.active
    ws.title = "Tổng hợp chi bộ"
    cols = ["STT", "Chi bộ (theo Excel)", "Thư mục chi bộ ghép được", "Số người theo Excel", "Có thư mục",
            "Khớp", "Cần sửa", "Chưa có thư mục", "Ngoài danh sách", "Hồ sơ đủ", "Hồ sơ thiếu"]
    n = len(cols)
    _title(ws, "BÁO CÁO ĐỐI CHIẾU DANH SÁCH HỆ THỐNG VỚI THƯ MỤC HỒ SƠ", sub, n)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n)
    c = ws.cell(row=3, column=1, value=(
        f"Excel: {len(cmp.people)} người · Có thư mục: {sum(1 for r in cmp.rows if r.person and r.dossier)} · "
        f"Khớp: {cmp.count(C.K_MATCH)} · Cần sửa: {sum(cmp.count(k) for k in C.NEED_FIX)} · "
        f"Chưa có thư mục: {cmp.count(C.K_NO_FOLDER)} · Ngoài danh sách: {cmp.count(C.K_NOT_IN_LIST)}"))
    c.font = _font(True, "1E3A8A", 12)
    c.alignment = CENTER
    head = 5
    fills = {6: PatternFill("solid", fgColor=GREEN), 7: PatternFill("solid", fgColor="A16207"),
             8: PatternFill("solid", fgColor=RED), 9: PatternFill("solid", fgColor=RED)}
    _header(ws, head, cols, [6, 36, 32, 11, 10, 9, 9, 11, 11, 9, 9], fills)
    r = head + 1
    for i, g in enumerate(C.summary_by_chibo(cmp), 1):
        r = _row(ws, r, [i, g["chibo"], g["folder"] or "(chưa có)", g["excel"], g["has_folder"], g["match"], g["fix"],
                         g["no_folder"], g["not_in_list"], g["full"], g["lack"]],
                 [CENTER, WRAP, WRAP] + [CENTER] * (n - 3))
        for col, key, color in ((7, "fix", AMBER), (8, "no_folder", RED), (9, "not_in_list", RED), (11, "lack", RED)):
            if g[key]:
                ws.cell(row=r - 1, column=col).font = _font(True, color)
        if not g["folder"]:
            ws.cell(row=r - 1, column=3).fill = FILL_NOTE
    last = r - 1
    if last > head:
        tot = ["", "TỔNG CỘNG", ""] + [f"=SUM({chr(64 + c)}{head + 1}:{chr(64 + c)}{last})" for c in range(4, n + 1)]
        _row(ws, r, tot, [CENTER, WRAP, WRAP] + [CENTER] * (n - 3))
        for cidx in range(1, n + 1):
            ws.cell(row=r, column=cidx).fill = FILL_TOTAL
            ws.cell(row=r, column=cidx).font = _font(True)
    _setup(ws, head, last, n)
    r += 2
    for t in [f"Bỏ qua {len(cmp.skipped)} dòng Excel không có CCCD/họ tên hợp lệ: "
              + "; ".join(f"dòng {a} ({b})" for a, b in cmp.skipped[:8])] if cmp.skipped else []:
        ws.cell(row=r, column=2, value="• " + t).font = _font(False, AMBER, 11)

    # ---- 2…5. danh sách
    _detail_sheet(wb, cmp, "ĐỐI CHIẾU CHI TIẾT TỪNG NGƯỜI", "Đối chiếu chi tiết", cmp.rows, sub)
    _detail_sheet(wb, cmp, "CÓ TRONG DANH SÁCH NHƯNG CHƯA CÓ THƯ MỤC HỒ SƠ", "Chưa có thư mục",
                  [r for r in cmp.rows if r.kind == C.K_NO_FOLDER], sub, check_col=True)
    _detail_sheet(wb, cmp, "THƯ MỤC HỒ SƠ KHÔNG CÓ TRONG DANH SÁCH", "Ngoài danh sách",
                  [r for r in cmp.rows if r.kind == C.K_NOT_IN_LIST], sub, check_col=True)
    _detail_sheet(wb, cmp, "CẦN SỬA: LỆCH TÊN · KHÁC CHI BỘ · NGHI GÕ NHẦM CCCD · CÓ GHI CHÚ", "Cần sửa",
                  [r for r in cmp.rows if r.kind in C.NEED_FIX or (r.kind == C.K_MATCH and r.notes)], sub,
                  check_col=True)

    # ---- 6. Giải thích
    ws = wb.create_sheet("Giải thích")
    _title(ws, "GIẢI THÍCH BÁO CÁO ĐỐI CHIẾU", "So khớp theo số CCCD (khoá chính), sau đó so họ tên và chi bộ", 2)
    items = [
        (C.K_MATCH, "Có trong danh sách và có thư mục; họ tên (bỏ dấu) giống; chi bộ khớp"),
        (C.K_NAME, "Số CCCD khớp nhưng họ tên trong tên thư mục khác danh sách → đổi tên thư mục theo danh sách"),
        (C.K_CHIBO, "Số CCCD và họ tên khớp nhưng thư mục đang nằm ở chi bộ khác → chuyển sang đúng chi bộ"),
        (C.K_TYPO, "Không trùng số CCCD nhưng lệch 1 chữ số hoặc đảo 2 chữ số liền nhau, họ tên giống "
                   "→ kiểm tra lại số CCCD (thường gõ nhầm khi đặt tên thư mục)"),
        (C.K_NO_FOLDER, "Có trong danh sách hệ thống nhưng chưa có thư mục hồ sơ → cần số hoá / tạo thư mục"),
        (C.K_NOT_IN_LIST, "Có thư mục nhưng không có trong danh sách → sai CCCD, đã chuyển đi hoặc thư mục thừa"),
        ("Hồ sơ / 01…05", "Tình trạng file bắt buộc lấy từ báo cáo hồ sơ: Đủ / Thiếu (Còn thiếu ghi mã)"),
        ("Ghi chú", "Trùng CCCD, CCCD Excel bị mất số 0 đầu (đã tự thêm), CCCD gần giống nhưng tên khác, "
                    "tên thư mục chưa đúng chuẩn…"),
        ("Chi bộ", "Ghép gần đúng tên chi bộ trong Excel với thư mục chi bộ (bỏ dấu, bỏ chữ “Chi bộ”)"),
    ]
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    r = 4
    from .report_excel import BORDER
    for k, v in items:
        for col, val, f in ((1, k, _font(True, "1E3A8A")), (2, v, _font())):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = f
            cell.alignment = WRAP
            cell.border = BORDER
        ws.row_dimensions[r].height = 34
        r += 1
    _setup(ws, 3, r - 1, 2, landscape=False, filter_=False)
    ws.freeze_panes = None

    wb.properties.title = "Đối chiếu danh sách hồ sơ"
    wb.properties.creator = "PDF-Read Pro"
    tmp = path + ".partial"
    wb.save(tmp)
    os.replace(tmp, path)
    return path
