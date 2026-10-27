"""Xuất Excel kết quả Gom ký số: tổng hợp theo hồ sơ (đã ký đủ / còn chờ / lỗi) + chi tiết từng file."""
from __future__ import annotations

import os
from datetime import datetime
from typing import Sequence

from openpyxl import Workbook

from . import signgather as G
from .report_excel import (CENTER, FILL_NO, FILL_NOTE, FILL_OK, FILL_TOTAL, WRAP, _font, _header, _row, _setup,
                           _title)

DONE = {G.R_SIGNED, G.R_HAVE}                              # bản ký hợp lệ (sẽ trả / đã có ở đích)


def export(path: str, rows: Sequence[tuple], title: str, sub: str) -> str:
    """rows: (hồ sơ, tên file, đích, mã trạng thái, chữ trạng thái, ghi chú, là bản ký)."""
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    if not os.path.isdir(os.path.dirname(os.path.abspath(path))):
        raise OSError("Thư mục lưu không tồn tại.")
    sub = f"{sub}   ·   Ngày lập: {datetime.now():%d/%m/%Y %H:%M}"
    wb = Workbook()

    # --- tổng hợp theo hồ sơ
    ws = wb.active
    ws.title = "Theo hồ sơ"
    per: dict = {}
    for folder, _, _, st, _, _, signed in rows:
        d = per.setdefault(folder or "?", {"ok": 0, "wait": 0, "bad": 0, "files": 0})
        d["files"] += 1
        if signed and st in DONE:
            d["ok"] += 1
        elif st in (G.R_WAIT, G.R_BUSY):
            d["wait"] += 1
        elif st in G.ISSUES:
            d["bad"] += 1
    cols = ["STT", "Hồ sơ (CCCD_HoVaTen)", "Số dòng", "Bản ký hợp lệ", "Còn chờ ký", "Lỗi / cần xử lý", "Kết luận"]
    _title(ws, title, sub, len(cols))
    _header(ws, 3, cols, [7, 42, 10, 13, 12, 14, 30])
    r = 4
    tot = [0, 0, 0, 0]
    for i, (folder, d) in enumerate(sorted(per.items()), 1):
        verdict = "Cần xử lý" if d["bad"] else ("Chờ ký" if d["wait"] else ("Đã ký" if d["ok"] else "—"))
        r = _row(ws, r, [i, folder, d["files"], d["ok"], d["wait"], d["bad"], verdict],
                 [CENTER, WRAP, CENTER, CENTER, CENTER, CENTER, WRAP])
        ws.cell(row=r - 1, column=7).fill = FILL_NO if d["bad"] else (FILL_NOTE if d["wait"] else FILL_OK)
        for k, key in enumerate(("files", "ok", "wait", "bad")):
            tot[k] += d[key]
    r = _row(ws, r, ["", "TỔNG", *tot, f"{len(per)} hồ sơ"], [CENTER, WRAP, CENTER, CENTER, CENTER, CENTER, WRAP])
    for c in range(1, len(cols) + 1):
        ws.cell(row=r - 1, column=c).fill = FILL_TOTAL
        ws.cell(row=r - 1, column=c).font = _font(True)
    _setup(ws, 3, r - 2, len(cols), landscape=False)

    # --- chi tiết
    ws = wb.create_sheet("Chi tiết")
    cols = ["STT", "Hồ sơ", "Tên file", "Đích", "Trạng thái", "Ghi chú"]
    _title(ws, title + " – chi tiết", sub, len(cols))
    _header(ws, 3, cols, [7, 32, 52, 52, 36, 40])
    r = 4
    for i, (folder, name, dest, st, text, note, _) in enumerate(rows, 1):
        r = _row(ws, r, [i, folder, name, dest, text, note], [CENTER, WRAP, WRAP, WRAP, WRAP, WRAP])
        if st in G.ISSUES:
            ws.cell(row=r - 1, column=5).fill = FILL_NO
        elif st in (G.R_WAIT, G.R_BUSY):
            ws.cell(row=r - 1, column=5).fill = FILL_NOTE
        elif st in DONE or st in G.COPY_ACTIONS or st == G.G_NEW:
            ws.cell(row=r - 1, column=5).fill = FILL_OK
    _setup(ws, 3, r - 1, len(cols))
    wb.properties.title = title
    wb.properties.creator = "PDF-Read Pro"
    tmp = path + ".partial"
    wb.save(tmp)
    os.replace(tmp, path)
    return path
